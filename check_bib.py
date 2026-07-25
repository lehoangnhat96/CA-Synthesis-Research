import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from docx import Document

docx_path = "D:/1 Master's Ana Chem/1 Master's thesis/Carbon Aerogel/Thesi grad/DCLV_CA_02.07 Fe-N CA_Renumbered.docx"
doc = Document(docx_path)

start = -1
end = -1
for i, p in enumerate(doc.paragraphs):
    if 'TÀI LIỆU THAM KHẢO' in p.text.upper():
        start = i
    if 'PHỤ LỤC' in p.text.upper() and start != -1:
        end = i
        break

print(f"Bibliography range: [{start}] to [{end}]")
for i in range(start, min(start+30, end)):
    p = doc.paragraphs[i]
    print(f"[{i}] Style={repr(p.style.name)} | text={repr(p.text[:90])}")
