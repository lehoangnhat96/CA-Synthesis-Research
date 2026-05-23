# -*- coding: utf-8 -*-
"""
Generate OUTLINE THÔNG SỐ TỔNG HỢP CARBON AEROGEL (N-CA và Fe/N-CA)
as a .docx file with formatted tables.
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import os

OUTPUT = os.path.join(
    r"D:\1 Thạc sĩ Hóa Phân Tích\ĐỀ TÀI_LVTN\0 Quy trình",
    "OUTLINE_THONG_SO_TONG_HOP_CA.docx",
)

doc = Document()

# ── Global styles ──
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.15

# Heading styles
for level in range(1, 4):
    hs = doc.styles[f"Heading {level}"]
    hs.font.name = "Times New Roman"
    hs.font.bold = True
    hs.font.color.rgb = RGBColor(0, 51, 102)
    if level == 1:
        hs.font.size = Pt(16)
    elif level == 2:
        hs.font.size = Pt(13)
    else:
        hs.font.size = Pt(12)


# ── Helper functions ──
def add_table(doc, headers, rows, col_widths=None):
    """Add a formatted table."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row shading
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    for r_idx, row_data in enumerate(rows):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10)
            # Alternate row shading
            if r_idx % 2 == 0:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F0FE"/>')
                cell._tc.get_or_add_tcPr().append(shading)

    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)

    doc.add_paragraph()  # spacer
    return table


def add_note(doc, text, prefix="LƯU Ý"):
    """Add a highlighted note paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(f"⚠ {prefix}: ")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 60, 0)
    run.font.size = Pt(10)
    run2 = p.add_run(text)
    run2.font.size = Pt(10)
    run2.font.italic = True


def add_bullet(doc, text, level=0):
    """Add a bullet point."""
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.27 + level * 0.63)
    for run in p.runs:
        run.font.size = Pt(11)


# ═══════════════════════════════════════════════════════
#  TITLE PAGE
# ═══════════════════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_t = p_title.add_run("OUTLINE THÔNG SỐ TỔNG HỢP\nCARBON AEROGEL (N-CA VÀ Fe/N-CA)")
run_t.bold = True
run_t.font.size = Pt(22)
run_t.font.color.rgb = RGBColor(0, 51, 102)

doc.add_paragraph()

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_s = p_sub.add_run("Đối chiếu — Xác thực kế thừa — Đánh giá từ Thư viện Tài liệu")
run_s.font.size = Pt(14)
run_s.font.color.rgb = RGBColor(80, 80, 80)

doc.add_paragraph()
p_info = doc.add_paragraph()
p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_i = p_info.add_run("Nguồn: Xơ dừa Bến Tre | IUH TPHCM | Hóa Phân tích\nSample set: N-CA (700°C) và Fe/N-CA (800°C)")
run_i.font.size = Pt(12)
run_i.font.italic = True

doc.add_page_break()

# ═══════════════════════════════════════════════════════
#  PHẦN 1: ĐỐI CHIẾU TÀI LIỆU
# ═══════════════════════════════════════════════════════
doc.add_heading("PHẦN 1. ĐỐI CHIẾU TÀI LIỆU VỚI 8 GIAI ĐOẠN TỔNG HỢP", level=1)

p = doc.add_paragraph()
run = p.add_run("Thư mục '1 Nguồn tài liệu' chứa ~90 bài báo PDF. Dưới đây là kết quả rà soát, ánh xạ vào từng giai đoạn và đánh giá % kế thừa.")
run.font.size = Pt(11)
run.font.italic = True

# ── SUMMARY TABLE ──
doc.add_heading("Tổng hợp đánh giá kế thừa toàn quy trình", level=2)

add_table(doc,
    ["Giai đoạn", "% Kế thừa", "Số bài hỗ trợ", "Thiếu sót chính"],
    [
        ["1. Tiền xử lý kiềm", "~90%", "17", "Gravimetry xơ dừa Bến Tre (original)"],
        ["2. Sol-gel & siêu âm", "~95%", "7", "(Gần hoàn chỉnh)"],
        ["3. Đúc khuôn & gel hóa", "~70%", "4", "Loại khuôn, thời gian ổn định"],
        ["4. Trao đổi DM & sấy thăng hoa", "~85%", "7", "Áp suất buồng sấy, tỷ lệ co rút"],
        ["5. Nhiệt phân", "~95%", "9", "(Gần hoàn chỉnh)"],
        ["6. Doping Fe", "~80%", "7", "Thời gian ngâm 24h vs 12h"],
        ["7. Hậu xử lý", "~75%", "4", "HCl vs H₂SO₄; T_anneal = 750°C"],
        ["8. Chế tạo điện cực & đo", "~90%", "22", "Tối ưu loading, chitosan vs nafion"],
        ["TRUNG BÌNH", "~85%", "~77 bài", ""],
    ],
    col_widths=[5, 2.5, 2.5, 7],
)

# ── GĐ 1 ──
doc.add_heading("GĐ 1 — Tiền xử lý kiềm (Alkali Pretreatment)", level=2)
doc.add_heading("Tài liệu tham khảo ánh xạ", level=3)

add_table(doc,
    ["#", "Tên tài liệu (PDF)", "Thông số được xác thực"],
    [
        ["1", "Cellulose Extraction from Coconut Coir with Alkaline Delignification Process", "NaOH 4–6%, 80°C, cơ chế loại lignin"],
        ["2", "Effect of NaOH Treatment on Coconut Coir Fibre…", "Ảnh hưởng nồng độ NaOH (2–10%) lên hình thái sợi dừa"],
        ["3", "Effects of Liq-to-Solid Ratio and Reaction Temp…", "Tỷ lệ lỏng:rắn, nhiệt độ tối ưu"],
        ["4", "Multi-Objective Optimization… Coir Fiber from Coconut", "Tối ưu đa mục tiêu cơ tính sợi dừa sau xử lý"],
        ["5", "Effect of Bleaching on Physicochemical…", "So sánh có/không tẩy trắng H₂O₂"],
        ["6", "Investigate coir fibers… Ben Tre… NaOH", "Xơ dừa Bến Tre — xử lý NaOH"],
        ["7", "Green Fabrication of Bio-based Aerogels from Coconut Coir", "Quy trình tạo aerogel từ xơ dừa"],
    ],
    col_widths=[1, 9, 7],
)

doc.add_heading("Đánh giá mức độ kế thừa", level=3)
add_table(doc,
    ["Thông số", "Xác thực?", "Ghi chú"],
    [
        ["NaOH 5–6%", "✅ Có", "Refs #1, #2, #3"],
        ["Nhiệt độ 75–80°C", "✅ Có", "Refs #1, #3, #4"],
        ["Thời gian 2–4h", "✅ Có", "Refs #1, #5"],
        ["Tỷ lệ rắn:lỏng", "✅ Có", "Ref #3"],
        ["Không tẩy trắng (no bleaching)", "✅ Có", "Ref #5 (so sánh có/không)"],
        ["Weight-loss xơ dừa Bến Tre", "⚠️ Thiếu", "Chưa có benchmark → Original Contribution"],
    ],
    col_widths=[5, 2, 10],
)

p_k = doc.add_paragraph()
run_k = p_k.add_run("→ Kế thừa: ~90%. ")
run_k.bold = True
run_k.font.color.rgb = RGBColor(0, 100, 0)
p_k.add_run("Thiếu sót duy nhất: chưa có dữ liệu gravimetry riêng cho xơ dừa Bến Tre (sẽ là Original Contribution).")

# ── GĐ 2 ──
doc.add_heading("GĐ 2 — Sol-gel & đánh siêu âm", level=2)
doc.add_heading("Tài liệu tham khảo ánh xạ", level=3)

add_table(doc,
    ["#", "Tên tài liệu (PDF)", "Thông số được xác thực"],
    [
        ["1", "Nitrogen-Doped CA… Coir Fibers… Ammonia-Urea System (BÀI CỐT LÕI)", "Hệ NH₄OH:Urea:H₂O 11:4:5, tỷ lệ 1:16, siêu âm, 0–5°C"],
        ["2", "Synthesis of Cellulose Aerogels from Coir via NaOH…", "Hệ NaOH:Urea — đối chứng (đã loại bỏ)"],
        ["3", "New Insights on the Role of Urea… Cellulose in Aqueous Alkali", "Vai trò Urea trong hòa tan cellulose"],
        ["4", "Unique Gelation Behavior of Cellulose in NaOH-Urea…", "Động học gel hóa cellulose"],
        ["5", "Controlling N-Doping Nature at Carbon Aerogels from Biomass…", "Kiểm soát dạng N-doping (pyridinic vs graphitic)"],
    ],
    col_widths=[1, 9, 7],
)

doc.add_heading("Đánh giá mức độ kế thừa", level=3)
add_table(doc,
    ["Thông số", "Xác thực?", "Ghi chú"],
    [
        ["Hệ NH₄OH:Urea:H₂O (11:4:5)", "✅ Có", "Fauziyah 2020 (doi:10.1021/acs.iecr.0c03771)"],
        ["Tỷ lệ rắn/lỏng 1:16", "✅ Có", "Ref #1"],
        ["Thứ tự: NH₃ → Urea → H₂O → Cellulose", "✅ Có", "Ref #1"],
        ["Siêu âm 30 phút, pulse 2s/1s, ice bath ≤10°C", "✅ Có", "Ref #1"],
        ["KHÔNG dùng NaOH:Urea", "✅ Có", "Ref #2 làm đối chứng loại bỏ"],
    ],
    col_widths=[5.5, 2, 9.5],
)

p_k2 = doc.add_paragraph()
run_k2 = p_k2.add_run("→ Kế thừa: ~95%. ")
run_k2.bold = True
run_k2.font.color.rgb = RGBColor(0, 100, 0)
p_k2.add_run("Fauziyah 2020 cung cấp gần như toàn bộ thông số sol-gel.")

# ── GĐ 3–7 (compact format) ──
for gd_num, gd_title, pct, refs_count, gaps in [
    (3, "Đúc khuôn & gel hóa", "~70%", 4, "Loại khuôn (silicone vs PP) và thời gian ổn định 30–60 phút chủ yếu là tối ưu nội bộ. Đề xuất: ghi nhận trong lab notebook."),
    (4, "Trao đổi dung môi & sấy thăng hoa", "~85%", 7, "Áp suất buồng sấy phụ thuộc thiết bị lab. Tỷ lệ co rút (shrinkage %) cần đo và so sánh. Đề xuất: đo kích thước gel trước/sau sấy."),
    (5, "Nhiệt phân (Pyrolysis)", "~95%", 9, "Chương trình nhiệt 3 ramp được xác thực rõ ràng. Mốc 700°C và 800°C được lý giải đầy đủ. Gần hoàn chỉnh."),
    (6, "Doping Fe (Post-impregnation)", "~80%", 7, "Thời gian ngâm 24h (so với literature 12h) cần biện luận: aerogel có macropore → cần thời gian dài hơn. Đề xuất: khảo sát 12h vs 24h, so sánh XPS Fe at%."),
    (7, "Hậu xử lý (Acid Leaching + Annealing)", "~75%", 4, "HCl vs H₂SO₄: cần thống nhất (đề xuất HCl 0.5M). T_anneal = 750°C nằm trong khoảng 700–800°C literature, cần biện luận thêm."),
]:
    doc.add_heading(f"GĐ {gd_num} — {gd_title}", level=2)
    p = doc.add_paragraph()
    run_p = p.add_run(f"Kế thừa: {pct} — {refs_count} bài hỗ trợ")
    run_p.bold = True
    if "95" in pct or "90" in pct or "85" in pct:
        run_p.font.color.rgb = RGBColor(0, 100, 0)
    elif "80" in pct:
        run_p.font.color.rgb = RGBColor(180, 130, 0)
    else:
        run_p.font.color.rgb = RGBColor(180, 60, 0)
    doc.add_paragraph(gaps)

# ── GĐ 8 ──
doc.add_heading("GĐ 8 — Chế tạo điện cực & đặc trưng điện hóa", level=2)

p_k8 = doc.add_paragraph()
run_k8 = p_k8.add_run("→ Kế thừa: ~90% — 22 bài hỗ trợ")
run_k8.bold = True
run_k8.font.color.rgb = RGBColor(0, 100, 0)

doc.add_paragraph("Bao gồm cả textbook chuẩn (Bard & Faulkner), review cảm biến điện hóa, và các bài về setup GCE, drop-casting, binder (chitosan/Nafion).")
doc.add_paragraph("Thiếu sót: Tối ưu loading drop-cast (khảo sát 3, 5, 7, 10 µL), so sánh chitosan vs Nafion binder.")

# ── TỔNG HỢP THIẾU SÓT ──
doc.add_heading("Tổng hợp thiếu sót & đề xuất bù đắp", level=2)

add_table(doc,
    ["Thiếu sót", "Giai đoạn", "Mức độ", "Đề xuất bù đắp"],
    [
        ["Gravimetry xơ dừa Bến Tre", "GĐ 1", "Thấp", "Original Contribution — đo và báo cáo"],
        ["Loại khuôn đúc", "GĐ 3", "Thấp", "Ghi nhận trong Experimental"],
        ["Áp suất buồng sấy thăng hoa", "GĐ 4", "TB", "Ghi thông số thiết bị cụ thể"],
        ["Tỷ lệ co rút (shrinkage %)", "GĐ 4", "TB", "Đo kích thước gel trước/sau sấy"],
        ["Thời gian ngâm Fe (24h vs 12h)", "GĐ 6", "TB", "Khảo sát 12h vs 24h, so XPS Fe at%"],
        ["HCl vs H₂SO₄ acid leaching", "GĐ 7", "TB", "Thống nhất: HCl 0.5M"],
        ["T_anneal = 750°C", "GĐ 7", "TB", "Biện luận: thấp hơn T_pyrolysis → tránh phá xốp"],
        ["Tối ưu loading drop-cast", "GĐ 8", "TB", "Khảo sát 3, 5, 7, 10 µL"],
        ["Chitosan vs Nafion binder", "GĐ 8", "Thấp", "Thử cả 2, chọn dựa trên ΔEp và RSD"],
    ],
    col_widths=[5, 2, 1.5, 8.5],
)

doc.add_page_break()

# ═══════════════════════════════════════════════════════
#  PHẦN 2: OUTLINE HOÀN CHỈNH
# ═══════════════════════════════════════════════════════
doc.add_heading("PHẦN 2. OUTLINE THÔNG SỐ TỔNG HỢP HOÀN CHỈNH", level=1)

p_note = doc.add_paragraph()
run_note = p_note.add_run("Tập trung trình bày: Thông số / Điều kiện  ↔  Mục đích & Lưu ý")
run_note.bold = True
run_note.font.size = Pt(12)
run_note.font.color.rgb = RGBColor(0, 51, 102)

# ── GĐ 1 FULL ──
doc.add_heading("GĐ 1 — Tiền xử lý kiềm (Alkali Pretreatment)", level=2)
doc.add_paragraph("Mục đích: Loại lignin/hemicellulose, lộ cellulose, giữ cấu trúc phân cấp tự nhiên.").bold = False

add_table(doc,
    ["Thông số", "Điều kiện", "Mục đích & Lưu ý"],
    [
        ["Nguyên liệu", "40g xơ dừa Bến Tre (đã xay nhuyễn)", "Xay trước để tăng bề mặt tiếp xúc kiềm"],
        ["Rửa sơ bộ", "800 mL nước vòi, 80 ± 2°C, khuấy 610 rpm, 4 giờ", "Loại bụi bẩn, tạp cơ học. Chắt bỏ nước sau lắng"],
        ["Kiềm hóa NaOH", "6% w/v (30g NaOH / 500 mL), 80 ± 2°C, 4 giờ, 610 rpm", "Phá vỡ liên kết lignin-cellulose"],
        ["Lọc & rửa", "Giấy lọc, rửa nước máy đến pH 6–7", "Kiểm tra pH bằng giấy chỉ thị"],
        ["Sấy khô", "80°C / 4 giờ", "Dàn đều tránh vón. Nếu vón → sấy lại + nghiền"],
        ["Tẩy trắng (tuỳ chọn)", "10% H₂O₂ + 3% NaOH (1:1), 1g:20mL, 80°C, 1h × 2", "Loại lignin tàn dư. Có thể bỏ qua"],
        ["Cân khối lượng", "Sấy 105°C, cân trước/sau kiềm hóa", "Dữ liệu gốc: tính % weight-loss"],
    ],
    col_widths=[3.5, 7, 6.5],
)

add_note(doc, "Trong quá trình đun có thể thất thoát nước → bù thêm nước cất (không dùng nước vòi).", "LƯU Ý")

# ── GĐ 2 FULL ──
doc.add_heading("GĐ 2 — Tổng hợp Sol-gel & đánh siêu âm", level=2)
doc.add_paragraph("Mục đích: Hòa tan cellulose tạo Sol đồng nhất, đưa N (từ NH₃ + Urea) vào mạng lưới gel.")

doc.add_heading("Công thức hệ dung môi (cho 1g cellulose)", level=3)
add_table(doc,
    ["Thành phần", "Lượng", "Vai trò"],
    [
        ["NH₃ 25%", "11 mL", "Tạo kiềm hòa tan cellulose + nguồn N-doping"],
        ["Urea", "4 g", "Phá cấu trúc hydrat hóa cellulose + nguồn N-doping"],
        ["H₂O DI", "5 mL", "Dung môi pha loãng"],
        ["Cellulose pulp", "1 g", "—"],
        ["Tổng dung môi", "16 mL", "Tỷ lệ rắn/lỏng = 1:16"],
    ],
    col_widths=[3.5, 3, 10.5],
)

doc.add_heading("Quy trình thực hiện", level=3)
add_table(doc,
    ["Bước", "Điều kiện", "Mục đích & Lưu ý"],
    [
        ["Pha dung môi", "NH₃ → Urea (khuấy tan) → H₂O", "TUYỆT ĐỐI KHÔNG dùng NaOH thay NH₃"],
        ["Cho cellulose", "Cho vào sau cùng, bọc kín màng thực phẩm", "Tránh NH₃ bay hơi"],
        ["Đánh siêu âm", "30 phút, pulse 2s on / 1s off", "Phân tán đồng nhất cellulose"],
        ["Kiểm soát nhiệt", "Ice bath ≤ 10°C (ưu tiên 0–5°C)", "BẮT BUỘC — NH₃ bay hơi nếu > 10°C"],
        ["Dấu hiệu đạt", "Sol đồng nhất, không còn sợi xơ thô", "Nếu chưa đạt → tăng thời gian siêu âm"],
        ["Bảo quản sol", "-12°C / 24h trong tủ lạnh", "Ổn định mạng lưới trước gel hóa"],
    ],
    col_widths=[3.5, 6, 7.5],
)

add_note(doc, "Nếu NH₃ bay hơi khi siêu âm → nồng độ N-doping giảm → XPS N at% < 3% → tín hiệu điện hóa yếu. Luôn giữ ice bath và bọc kín.", "CẢNH BÁO")

# ── GĐ 3 FULL ──
doc.add_heading("GĐ 3 — Đúc khuôn & gel hóa (Casting & Gelation)", level=2)
doc.add_paragraph("Mục đích: Chuyển Sol thành cấu trúc Hydrogel 3D ổn định.")

add_table(doc,
    ["Bước", "Điều kiện", "Mục đích & Lưu ý"],
    [
        ["Đổ khuôn", "Khuôn silicone hoặc ống PP (KHÔNG dùng thủy tinh)", "Thủy tinh dính gel, khó tách"],
        ["Nhiệt độ khuôn", "0–5°C", "Giữ lạnh tránh gel hóa bất thường"],
        ["Ổn định", "Để yên 30–60 phút", "Mạng lưới hình thành ban đầu"],
        ["Cấp đông gel hóa", "-14°C đến -20°C / 24 giờ", "Tạo mạng hydrogel tổ ong"],
        ["Rã đông", "Tự nhiên 10–20 phút", "Loại bỏ nước dư bề mặt"],
    ],
    col_widths=[3.5, 7, 6.5],
)

# ── GĐ 4 FULL ──
doc.add_heading("GĐ 4 — Keo tụ, trao đổi dung môi & sấy thăng hoa", level=2)
doc.add_paragraph("Mục đích: Thay nước trong gel bằng dung môi dễ thăng hoa, sau đó sấy loại dung môi mà không phá vỡ mạng 3D.")

doc.add_heading("4a. Keo tụ / Coagulation (chia 2 mẫu song song)", level=3)
add_table(doc,
    ["Mẫu", "Điều kiện ngâm", "Mục đích"],
    [
        ["N-CA", "Ethanol 98% lạnh (0–5°C), 15 mL/mL gel, 24–48h", "Thay thế nước bằng ethanol"],
        ["Fe/N-CA", "Ethanol + FeCl₃·6H₂O (0–5°C), 24–48h\n1% Fe → 0.048g; 5% Fe → 0.242g", "Keo tụ + tẩm Fe vào cấu trúc gel"],
    ],
    col_widths=[3, 9, 5],
)

doc.add_heading("4b. Rửa / Solvent Exchange", level=3)
add_table(doc,
    ["Bước", "Điều kiện", "Lưu ý"],
    [
        ["Rửa ethanol", "DI water ≥ 5× thể tích gel, ngâm 2–3h, lặp 2–3 lần", "Đến khi hết mùi cồn + pH trung tính"],
        ["Gradient TBA (optional)", "15% TBA 6–8h → 30% TBA 12–24h", "Tinh thể hóa tốt hơn khi cấp đông"],
    ],
    col_widths=[4, 8, 5],
)

add_note(doc, "KHÔNG rửa quá nhiều lần — tránh rửa trôi urea liên kết trong gel (giảm nguồn N-doping).")

doc.add_heading("4c. Sấy thăng hoa / Freeze Drying", level=3)
add_table(doc,
    ["Thông số", "Điều kiện", "Mục đích"],
    [
        ["Cấp đông", "-20°C đến -40°C, ≥ 12–24 giờ", "Đóng băng hoàn toàn dung môi"],
        ["Bẫy lạnh (Cold trap)", "≤ -30°C", "Gom hơi dung môi, bảo vệ bơm chân không"],
        ["Áp suất", "< 20 Pa (~0.15 Torr)", "Thăng hoa: rắn → khí, không qua lỏng"],
        ["Thời gian", "24–48 giờ (đến khi KL không đổi)", "—"],
    ],
    col_widths=[4, 6, 7],
)

add_note(doc, "Nếu sấy nhiệt thay vì thăng hoa → lực mao quản phá sập mạng gel → BET thấp, mesopore giảm → khuếch tán kém.", "CẢNH BÁO")

# ── GĐ 5 FULL ──
doc.add_heading("GĐ 5 — Nhiệt phân (Pyrolysis)", level=2)
doc.add_paragraph("Mục đích: Carbon hóa cellulose aerogel, hình thành cấu trúc carbon bán tinh thể với N-doping.")

doc.add_heading("Thiết lập khí bảo vệ", level=3)
add_table(doc,
    ["Bước", "Thông số", "Mục đích"],
    [
        ["Purge trước nung", "N₂, 150–200 mL/min, 15–30 phút", "Đẩy hết O₂ ra khỏi lò"],
        ["Duy trì khi nung", "N₂, 100 mL/min", "Bảo vệ suốt quá trình"],
        ["Duy trì khi cooling", "N₂, 100 mL/min", "Tránh oxy hóa khi carbon còn nóng"],
    ],
    col_widths=[4, 6, 7],
)

doc.add_heading("Chương trình nhiệt", level=3)
add_table(doc,
    ["Giai đoạn", "Khoảng nhiệt", "Tốc độ", "Giữ nhiệt", "Mục đích"],
    [
        ["Ramp 1", "25 → 150°C", "5°C/min", "30 phút", "Loại hoàn toàn ẩm"],
        ["Ramp 2", "150 → 400°C", "5°C/min", "30 phút", "Carbon hóa sơ bộ"],
        ["Ramp 3", "400°C → T_target", "5°C/min", "2 giờ", "Carbon hóa + N-doping"],
        ["Cooling", "T_target → < 50°C", "Tự nhiên", "—", "Giữ N₂ liên tục"],
    ],
    col_widths=[2.5, 3.5, 2.5, 2, 6.5],
)

doc.add_heading("Nhiệt độ mục tiêu", level=3)
add_table(doc,
    ["Mẫu", "T_target", "Lý do"],
    [
        ["N-CA", "700°C", "Giữ nhiều pyridinic-N. Đủ cho DPV paracetamol."],
        ["Fe/N-CA (nền)", "800°C", "Tăng graphitization → tăng độ dẫn → nền tốt cho Fe-Nₓ."],
    ],
    col_widths=[3, 2.5, 11.5],
)

# ── GĐ 6 FULL ──
doc.add_heading("GĐ 6 — Doping Fe (Post-impregnation) — CHỈ CHO Fe/N-CA", level=2)
doc.add_paragraph("Mục đích: Đưa Fe³⁺ vào mao quản carbon aerogel, tạo tiền chất cho tâm xúc tác Fe-Nₓ.")

add_table(doc,
    ["Thông số", "Điều kiện", "Mục đích & Lưu ý"],
    [
        ["Tiền chất", "FeCl₃·6H₂O (MW = 270.3)", "Tan tốt trong ethanol"],
        ["Dung môi", "Ethanol", "KHÔNG dùng nước (Fe(OH)₃ kết tủa trong kiềm)"],
        ["Nồng độ Fe", "1% (0.048g) và 5% (0.242g) / g mẫu", "Khảo sát tối ưu → chọn mẫu tốt nhất"],
        ["30 phút đầu", "Siêu âm nhẹ", "Đẩy Fe³⁺ vào mao quản ban đầu"],
        ["Thời gian còn lại", "Khuấy từ chậm, tổng 24h", "Cho Fe³⁺ khuếch tán sâu vào lõi aerogel"],
        ["Sấy sau tẩm", "60°C", "Loại ethanol nhẹ nhàng, giữ Fe phân tán"],
    ],
    col_widths=[3.5, 6, 7.5],
)

add_note(doc, "TUYỆT ĐỐI KHÔNG thêm Fe vào hệ NH₃/Urea/H₂O → gây kết tủa Fe(OH)₃ → phá hỏng sol-gel.", "CẢNH BÁO")

# ── GĐ 7 FULL ──
doc.add_heading("GĐ 7 — Hậu xử lý (Acid Leaching + Annealing) — CHỈ CHO Fe/N-CA", level=2)
doc.add_paragraph("Mục đích: Loại bỏ Fe tự do/Fe oxide không hoạt tính, ổn định tâm xúc tác Fe-Nₓ.")

doc.add_heading("7a. Acid Leaching", level=3)
add_table(doc,
    ["Thông số", "Điều kiện", "Mục đích"],
    [
        ["Acid", "HCl 0.5M", "Hòa tan Fe metallic, Fe₃C, Fe oxide"],
        ["Nhiệt độ", "80°C", "Tăng tốc phản ứng hòa tan"],
        ["Thời gian", "8 giờ", "Đủ để loại hết Fe không phối trí"],
        ["Sau xử lý", "Lọc → rửa DI water → sấy 60°C/12h", "Đưa về trạng thái khô sạch"],
    ],
    col_widths=[3.5, 6, 7.5],
)

doc.add_heading("7b. Annealing lần 2", level=3)
add_table(doc,
    ["Thông số", "Điều kiện", "Mục đích"],
    [
        ["Khí", "N₂", "Bảo vệ carbon"],
        ["Tốc độ gia nhiệt", "5°C/min", "—"],
        ["T_anneal", "750°C", "< T_pyrolysis (800°C) → tránh phá xốp. Đủ ổn định Fe-Nₓ"],
        ["Giữ nhiệt", "1 giờ", "—"],
        ["Cooling", "Tự nhiên trong N₂", "—"],
        ["Kiểm chứng ngay sau", "Raman I_D/I_G", "Xác nhận cấu trúc carbon không bị hỏng"],
    ],
    col_widths=[3.5, 5, 8.5],
)

# ── GĐ 8 FULL ──
doc.add_heading("GĐ 8 — Chế tạo điện cực & đặc trưng điện hóa", level=2)
doc.add_paragraph("Mục đích: Chuyển vật liệu bột CA thành điện cực làm việc, đánh giá hiệu năng phân tích.")

doc.add_heading("8a. Nghiền, rây & lưu trữ", level=3)
add_bullet(doc, "Nghiền trong cối mã não đến bột mịn đồng đều")
add_bullet(doc, "Rây qua lưới 75 µm (lấy phần < 75 µm)")
add_bullet(doc, "Bảo quản: Vial thủy tinh nâu, nắp kín, 4°C, hút chân không nhẹ hoặc bơm N₂")
add_bullet(doc, "Ghi nhãn đầy đủ: Mẫu / Batch / Ngày / BET / σ")

doc.add_heading("8b. Hoạt hóa GCE", level=3)
add_table(doc,
    ["Bước", "Điều kiện", "Mục đích"],
    [
        ["Mài", "Al₂O₃ 0.05 µm, hình số 8, 2 min × 2 lần", "Làm sạch + tạo bề mặt mới"],
        ["Siêu âm", "DI water 1 min × 3 + Ethanol 1 min × 1", "Loại bụi Al₂O₃"],
        ["Loại Al₂O₃ dư", "NaOH 0.1M + anodize +1.8V/10s", "—"],
        ["Kiểm tra", "CV 3 chu kỳ, 5 mM K₃Fe(CN)₆/0.1M KCl, 50 mV/s", "ΔEp < 70 mV; Ipa/Ipc = 0.95–1.05"],
    ],
    col_widths=[3, 8, 6],
)

doc.add_heading("8c. Pha Ink & Drop-casting", level=3)
add_table(doc,
    ["Thông số", "Điều kiện", "Lưu ý"],
    [
        ["Cân carbon", "5.0 mg ± 0.01 mg", "—"],
        ["Dung môi", "950 µL DMF (HPLC grade)", "Hoặc Nước:Ethanol"],
        ["Binder", "50 µL Nafion 0.25%/EtOH", "Hoặc Chitosan 1% trong acetic acid 1%"],
        ["Siêu âm", "Bath, nước đá (< 20°C), 30 phút", "KHÔNG probe > 30% amplitude"],
        ["Kiểm tra ink", "Đen đồng đều, lắng chậm", "DLS D₅₀ < 500 nm = tốt"],
        ["Drop-cast", "7 µL nhỏ chậm từng µL lên GCE", "Loading lý thuyết = 35 µg/điện cực"],
        ["Sấy màng", "Tự nhiên, 25°C, tủ hút sạch, 2–3h", "KHÔNG thổi N₂, KHÔNG sấy nóng"],
    ],
    col_widths=[3, 7, 7],
)

doc.add_heading("8d. Screening 4 vật liệu", level=3)
add_table(doc,
    ["Phép đo", "Dung dịch", "Thông số", "Tiêu chí Pass"],
    [
        ["CV", "5 mM K₃Fe(CN)₆ / 0.1M KCl", "50 mV/s, 5 chu kỳ", "Ip tăng ≥ 30%, ΔEp giảm"],
        ["EIS", "5 mM K₃Fe(CN)₆ / 0.1M KCl", "100 kHz → 0.1 Hz, 10 mV", "Rct giảm ≥ 30%"],
        ["ECSA", "0.1M H₂SO₄", "25–200 mV/s", "Tính C_dl → ECSA"],
        ["SWASV sơ bộ", "Pb²⁺+Cd²⁺ 100 ppb, acetate pH 4.5", "E_dep=-1.1V, t_dep=120s", "Tách peak ≥ 100 mV"],
    ],
    col_widths=[3, 5, 4.5, 4.5],
)

doc.add_heading("8e. Phân tích Pb²⁺/Cd²⁺/Zn²⁺ (SWASV) — Bài Q1/Q2", level=3)
add_table(doc,
    ["Thông số", "Khoảng khảo sát", "Giá trị khởi đầu"],
    [
        ["pH dung dịch nền", "3.5 – 6.0 (acetate buffer)", "4.5"],
        ["E_dep", "-0.8 đến -1.4 V vs Ag/AgCl", "-1.1 V"],
        ["t_dep", "60 – 240 s (khuấy 400 rpm)", "120 s"],
        ["SW frequency", "10 – 100 Hz", "25 Hz"],
        ["SW amplitude", "10 – 50 mV", "25 mV"],
        ["Step potential", "2 – 10 mV", "4–5 mV"],
    ],
    col_widths=[4, 7, 6],
)

add_bullet(doc, "LOD mục tiêu: Pb²⁺ < 0.1 µg/L; Zn²⁺ < 10 nM")
add_bullet(doc, "Đường chuẩn: 0.5, 1, 2, 5, 10, 20, 50, 100, 200 ppb (n=3/điểm)")
add_bullet(doc, "Ion gây nhiễu: Cu²⁺ ⚠️, Hg²⁺, Fe³⁺, Mn²⁺, Ca²⁺, Mg²⁺, Na⁺, K⁺, NO₃⁻")
add_bullet(doc, "Real sample: Nước sông/máy → lọc 0.45 µm → vô cơ hóa ướt → Standard Addition 3 điểm")

doc.add_heading("8f. Phân tích Paracetamol (DPV) — Bài Q3/Q4", level=3)
add_table(doc,
    ["Thông số", "Khoảng khảo sát", "Giá trị khởi đầu"],
    [
        ["pH dung dịch nền", "5.0 – 8.0 (PBS)", "7.0–7.4"],
        ["Pulse amplitude", "10 – 75 mV", "50 mV"],
        ["Pulse width", "10 – 100 ms", "50 ms"],
        ["Step potential", "2 – 10 mV", "5 mV"],
    ],
    col_widths=[4, 7, 6],
)

add_bullet(doc, "LOD mục tiêu: < 0.05 µM")
add_bullet(doc, "Chất gây nhiễu: AA, UA, DA, glucose, diclofenac, nimesulide, amoxicillin")
add_bullet(doc, "Real sample: Thuốc viên paracetamol thương mại → Recovery 95–105%")

doc.add_heading("8g. Repeatability, Reproducibility, Stability", level=3)
add_table(doc,
    ["Chỉ tiêu", "Phương pháp", "Tiêu chí"],
    [
        ["Repeatability", "1 điện cực × 10 lần liên tiếp", "RSD < 5%"],
        ["Reproducibility", "5 điện cực độc lập, cùng điều kiện", "RSD < 8%"],
        ["Stability", "Đo lại sau 0, 7, 14, 30 ngày (bảo quản 4°C)", "> 90% Ip sau 14 ngày"],
    ],
    col_widths=[3.5, 7, 6.5],
)

# ── QUALITY GATE ──
doc.add_page_break()
doc.add_heading("QUALITY GATE — CHECKLIST SAU MỖI GIAI ĐOẠN", level=2)

add_table(doc,
    ["Giai đoạn", "Dữ liệu bắt buộc", "Tiêu chí PASS", "Nếu FAIL"],
    [
        ["GĐ 1", "% weight-loss (sấy 105°C)", "Có dữ liệu trước/sau", "Thiếu Original Contribution"],
        ["GĐ 2", "Nhiệt độ ice bath suốt siêu âm", "Sol đồng nhất, T ≤ 10°C", "Tăng siêu âm, thêm đá"],
        ["GĐ 3", "Ảnh gel sau cấp đông", "Không nứt, không sụp", "Tối ưu ethanol exchange"],
        ["GĐ 4", "Hình dạng aerogel sau sấy", "Mạng 3D, co rút thấp", "Kiểm tra cấp đông"],
        ["GĐ 5", "Yield %, màu, độ giòn", "Đen, giòn, không cháy", "Kiểm tra purge N₂"],
        ["GĐ 6", "Lượng FeCl₃, thời gian ngâm", "—", "—"],
        ["GĐ 7", "Raman I_D/I_G sau annealing", "0.9–1.2", "Defect quá mức / thiếu sites"],
        ["GĐ 8 – Raman", "I_D/I_G", "0.9–1.2", "Điều chỉnh T nung"],
        ["GĐ 8 – XPS", "N at%, Fe at%", "N 3–6%, Fe 1–3%, N:Fe ≥ 4:1", "Điều chỉnh Fe / annealing"],
        ["GĐ 8 – BET", "S_BET", "> 300 m²/g", "KOH activation nếu cần"],
        ["GĐ 8 – CV", "ΔEp", "< 120 mV", "Tối ưu loading, binder"],
        ["GĐ 8 – EIS", "Rct", "Fe/N-CA < N-CA < GCE", "Kiểm tra Fe-Nₓ, màng"],
    ],
    col_widths=[2.5, 4.5, 4.5, 5.5],
)

p_final = doc.add_paragraph()
run_f = p_final.add_run("\n→ Nếu cả 5 chỉ số trên đều PASS → vật liệu đủ nền để chuyển sang tối ưu DPV/SWASV và xây dựng phần Analytical Performance cho bài báo.")
run_f.bold = True
run_f.font.color.rgb = RGBColor(0, 100, 0)
run_f.font.size = Pt(11)

# ── SAVE ──
doc.save(OUTPUT)
print(f"\n[OK] File Word da duoc tao thanh cong tai:\n   {OUTPUT}")
