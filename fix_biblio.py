from docx import Document
import re

docx_path = "D:/1 Master's Ana Chem/1 Master's thesis/Carbon Aerogel/Thesi grad/DCLV_CA_02.07 Fe-N CA_Renumbered.docx"
doc = Document(docx_path)

start_p = None
end_p = None
start_idx = -1
end_idx = -1

for i, p in enumerate(doc.paragraphs):
    if 'TÀI LIỆU THAM KHẢO' in p.text.upper():
        start_idx = i
        start_p = p
    elif 'PHỤ LỤC' in p.text.upper() and start_idx != -1:
        end_idx = i
        end_p = p
        break

if start_idx != -1 and end_idx != -1:
    print(f"Found biblio range: {start_idx} to {end_idx}")
    
    # Read the new bibliography
    with open("D:/1 Master's Ana Chem/1 Master's thesis/Carbon Aerogel/Thesi grad/DANH_MUC_MOI.txt", 'r', encoding='utf-8') as f:
        bib_lines = f.read().strip().split('\n')
        
    # Filter out empty lines
    bib_lines = [line.strip() for line in bib_lines if line.strip()]
    
    # Store the style of the first bibliography entry to copy it
    bib_style = None
    if end_idx > start_idx + 1:
        bib_style = doc.paragraphs[start_idx + 1].style
    
    # Clear the old bibliography paragraphs (we don't delete them, just clear text to avoid messing up internal XML structure)
    for i in range(start_idx + 1, end_idx):
        doc.paragraphs[i].text = ""
        
    # Insert new bibliography entries BEFORE the 'PHỤ LỤC' paragraph
    # We do it in reverse order so they appear correctly?
    # Wait, insert_paragraph_before inserts before the current element.
    # So if we iterate forward and insert before 'PHỤ LỤC', they will be in correct order!
    for line in bib_lines:
        new_p = end_p.insert_paragraph_before(line)
        if bib_style:
            new_p.style = bib_style

    doc.save(docx_path)
    print("Bibliography updated successfully!")
else:
    print("Could not find exact bounds for bibliography.")
