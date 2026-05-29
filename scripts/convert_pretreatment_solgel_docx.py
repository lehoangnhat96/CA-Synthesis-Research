# -*- coding: utf-8 -*-
"""
Dynamic Markdown to Word (.docx) converter for the Pretreatment & Sol-Gel SOP.
Preserves fully accented Vietnamese text, formatting, tables, and Unicode symbols
from "1 Phân Tích Chi Tiết Giai Đoạn #0 & #1_ Tiền Xử Lý và Sol-Gel.md".
"""

import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Ensure UTF-8 output encoding for console
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

# Define paths
WORKSPACE_DIR = Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
INPUT_MD = WORKSPACE_DIR / "1 Ref materials" / "09_Tai_lieu_Quy_trinh_AI" / "1 Phân Tích Chi Tiết Giai Đoạn #0 & #1_ Tiền Xử Lý và Sol-Gel.md"
OUTPUT_DOCX = WORKSPACE_DIR / "1 SOPs" / "new" / "1 Phân Tích Chi Tiết Giai Đoạn #0 & #1_ Tiền Xử Lý và Sol-Gel.docx"

def add_formatted_text(p, text):
    """Parses markdown bold (**) and adds formatted runs to a paragraph."""
    # Strip emojis or keep them based on preference (we'll keep them as they are in Unicode)
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part and idx == 0:
            continue
        run = p.add_run(part)
        if idx % 2 == 1:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102) # Dark blue for highlighted terms
        else:
            run.font.color.rgb = RGBColor(30, 30, 30) # Dark gray for normal text

def add_table_row(table, row_idx, row_data, is_header=False):
    """Fills a row in the docx table, applies shading and bold runs."""
    row = table.rows[row_idx]
    for col_idx, text in enumerate(row_data):
        if col_idx >= len(row.cells):
            break
        cell = row.cells[col_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        
        if is_header:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            # Deep blue header shading
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="003366"/>')
            cell._tc.get_or_add_tcPr().append(shading)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_formatted_text(p, text)
            p.runs[0].font.size = Pt(9.5)
            # Alternating light blue rows
            if row_idx % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F0FE"/>')
                cell._tc.get_or_add_tcPr().append(shading)

def add_note(doc, text, prefix="LUU Y"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(f"[!] {prefix}: ")
    run.bold = True
    run.font.color.rgb = RGBColor(180, 60, 0)
    run.font.size = Pt(10)
    run2 = p.add_run(text)
    run2.font.size = Pt(10)
    run2.font.italic = True

def parse_md_to_docx(md_path, docx_path):
    print(f"Đang đọc tệp tin: {md_path}")
    if not md_path.exists():
        print(f"LỖI: Không tìm thấy tệp {md_path}")
        return False
        
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    doc = Document()
    
    # -- Configure Normal Style --
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11.5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15
    
    # -- Document Cover / Header --
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    run_t = p_title.add_run("TÀI LIỆU CHUẨN HÓA QUY TRÌNH\nPHÂN TÍCH CHI TIẾT GIAI ĐOẠN #0 & #1")
    run_t.bold = True
    run_t.font.size = Pt(20)
    run_t.font.color.rgb = RGBColor(0, 51, 102)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_s = p_sub.add_run("TIỀN XỬ LÝ (DELIGNIFICATION) VÀ SOL-GEL Fe/N-CARBON AEROGEL\n(Nguồn xơ dừa Bến Tre - Chuẩn Unicode hóa học)")
    run_s.font.size = Pt(12)
    run_s.font.italic = True
    run_s.font.color.rgb = RGBColor(80, 80, 80)
    
    # Horizontal rule
    p_hr = doc.add_paragraph()
    p_hr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hr.paragraph_format.space_after = Pt(24)
    p_hr.add_run("__________________________________________________________________").font.color.rgb = RGBColor(180, 180, 180)
    
    table_accumulator = []
    in_table = False
    
    idx = 0
    while idx < len(lines):
        line = lines[idx].strip()
        
        # Detect Table
        if line.startswith("|"):
            in_table = True
            table_accumulator.append(line)
            idx += 1
            continue
        elif in_table:
            # Process and render accumulated table
            if table_accumulator:
                # Filter out separator row (e.g. |---|---|)
                clean_rows = []
                for row_str in table_accumulator:
                    cells = [c.strip() for c in row_str.split("|")[1:-1]]
                    if all(re.match(r"^[-:\s]+$", c) for c in cells) and len(cells) > 0:
                        continue # Skip separator
                    clean_rows.append(cells)
                
                if clean_rows:
                    headers = clean_rows[0]
                    data_rows = clean_rows[1:]
                    
                    # Create Word Table
                    t = doc.add_table(rows=len(clean_rows), cols=len(headers))
                    t.style = "Table Grid"
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    
                    # Header Row
                    add_table_row(t, 0, headers, is_header=True)
                    
                    # Data Rows
                    for r_idx, r_data in enumerate(data_rows):
                        add_table_row(t, r_idx + 1, r_data, is_header=False)
                    
                    doc.add_paragraph() # Add space after table
                    
            table_accumulator = []
            in_table = False
            
        if not line:
            idx += 1
            continue
            
        # Detect Heading 2 or Heading 1 based on syntax
        if line.startswith("## "):
            header_text = line[3:].replace("**", "").strip()
            if "GIAI ĐOẠN" in header_text:
                h = doc.add_heading(level=1)
                h.paragraph_format.space_before = Pt(18)
                h.paragraph_format.space_after = Pt(6)
                run = h.add_run(header_text)
                run.font.size = Pt(15)
                run.bold = True
                run.font.color.rgb = RGBColor(0, 51, 102)
            else:
                h = doc.add_heading(level=2)
                h.paragraph_format.space_before = Pt(12)
                h.paragraph_format.space_after = Pt(4)
                run = h.add_run(header_text)
                run.font.size = Pt(12.5)
                run.bold = True
                run.font.color.rgb = RGBColor(0, 51, 102)
                
        # Detect Bullet list
        elif line.startswith("- ") or line.startswith("• "):
            bullet_text = line[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(1.0)
            add_formatted_text(p, bullet_text)
            
        # Detect Warnings/Notes
        elif line.startswith("⚠") or line.startswith("[!]"):
            note_text = line.replace("⚠", "").replace("**", "").replace("Lưu ý:", "").strip()
            add_note(doc, note_text, "LƯU Ý QUAN TRỌNG")
            
        # Regular paragraph
        else:
            p = doc.add_paragraph()
            add_formatted_text(p, line)
            
        idx += 1
        
    # Make sure to flush last table if file ended on a table
    if in_table and table_accumulator:
        clean_rows = []
        for row_str in table_accumulator:
            cells = [c.strip() for c in row_str.split("|")[1:-1]]
            if all(re.match(r"^[-:\s]+$", c) for c in cells) and len(cells) > 0:
                continue
            clean_rows.append(cells)
        if clean_rows:
            headers = clean_rows[0]
            data_rows = clean_rows[1:]
            t = doc.add_table(rows=len(clean_rows), cols=len(headers))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            add_table_row(t, 0, headers, is_header=True)
            for r_idx, r_data in enumerate(data_rows):
                add_table_row(t, r_idx + 1, r_data, is_header=False)
            doc.add_paragraph()

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"[OK] Đã xuất thành công: {docx_path}")
    return True

if __name__ == "__main__":
    parse_md_to_docx(INPUT_MD, OUTPUT_DOCX)
