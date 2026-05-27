# -*- coding: utf-8 -*-
"""
High-fidelity converter for 2 Ref lib directory.
Recursively converts all unconverted .docx to .md, .xlsx to .csv, and .pdf to .md (with images)
within the D:\\1 Master's Ana Chem\\1 Master's thesis\\Carbon Aerogel\\2 Ref lib directory.
"""

import os
import sys
import csv
import shutil
import hashlib
from pathlib import Path

# Set console encoding to UTF-8 on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

# Import python-docx
try:
    from docx import Document
except ImportError:
    print("Error: 'python-docx' is not installed.")
    sys.exit(1)

# Import openpyxl
try:
    import openpyxl
except ImportError:
    print("Error: 'openpyxl' is not installed.")
    sys.exit(1)

# Import PyMuPDF
try:
    import fitz  # PyMuPDF
except ImportError:
    print("Warning: 'pymupdf' is not installed. PDF conversion will be skipped.")
    fitz = None

try:
    import pymupdf4llm
except ImportError:
    pymupdf4llm = None


def make_long_path(path):
    abs_path = os.path.abspath(path)
    if os.name == 'nt' and not abs_path.startswith('\\\\?\\'):
        abs_path = abs_path.replace('/', '\\')
        return '\\\\?\\' + abs_path
    return abs_path


def convert_docx_to_md(docx_path, md_path):
    print(f"Converting DOCX: {docx_path.name} -> {md_path.name}")
    try:
        doc = Document(make_long_path(docx_path))
    except Exception as e:
        print(f"  -> Failed to read DOCX {docx_path.name}: {e}")
        return False

    md_content = []
    body_elements = doc.element.body
    tables_iterator = iter(doc.tables)
    paragraphs_iterator = iter(doc.paragraphs)
    
    for child in body_elements:
        if child.tag.endswith('p'):
            try:
                p = next(paragraphs_iterator)
                text_runs = []
                for run in p.runs:
                    text = run.text
                    if not text:
                        continue
                    if run.bold:
                        text = f"**{text}**"
                    if run.italic:
                        text = f"*{text}*"
                    text_runs.append(text)
                
                paragraph_text = "".join(text_runs).strip()
                if not paragraph_text:
                    md_content.append("")
                    continue
                
                style_name = p.style.name.lower() if p.style and p.style.name else ""
                if "heading 1" in style_name:
                    md_content.append(f"# {paragraph_text}\n")
                elif "heading 2" in style_name:
                    md_content.append(f"## {paragraph_text}\n")
                elif "heading 3" in style_name:
                    md_content.append(f"### {paragraph_text}\n")
                elif "heading 4" in style_name:
                    md_content.append(f"#### {paragraph_text}\n")
                elif "list bullet" in style_name or paragraph_text.startswith(("- ", "* ", "• ")):
                    clean_text = paragraph_text
                    for prefix in ["- ", "* ", "• "]:
                        if clean_text.startswith(prefix):
                            clean_text = clean_text[len(prefix):]
                            break
                    md_content.append(f"- {clean_text}")
                else:
                    md_content.append(paragraph_text)
            except StopIteration:
                pass
        elif child.tag.endswith('tbl'):
            try:
                table = next(tables_iterator)
                table_lines = []
                if len(table.rows) > 0:
                    headers = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
                    table_lines.append("| " + " | ".join(headers) + " |")
                    table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    for row in table.rows[1:]:
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        table_lines.append("| " + " | ".join(cells) + " |")
                md_content.append("\n" + "\n".join(table_lines) + "\n")
            except StopIteration:
                pass

    try:
        with open(make_long_path(md_path), "w", encoding="utf-8") as f:
            f.write("\n\n".join(md_content))
        print(f"  -> Successfully saved MD: {md_path.name}")
        return True
    except Exception as e:
        print(f"  -> Failed to write MD {md_path.name}: {e}")
        return False


def convert_xlsx_to_csv(xlsx_path, csv_base_path):
    print(f"Converting XLSX: {xlsx_path.name}")
    try:
        wb = openpyxl.load_workbook(make_long_path(xlsx_path), data_only=True)
    except Exception as e:
        print(f"  -> Failed to read XLSX {xlsx_path.name}: {e}")
        return False

    success = True
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        if sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None:
            continue
            
        if len(wb.sheetnames) > 1:
            csv_path = csv_base_path.with_name(f"{csv_base_path.stem}_{sheet_name}.csv")
        else:
            csv_path = csv_base_path.with_name(f"{csv_base_path.stem}.csv")

        print(f"  -> Saving sheet '{sheet_name}' -> {csv_path.name}")
        try:
            with open(make_long_path(csv_path), "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                for r in range(1, sheet.max_row + 1):
                    row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
                    while row_vals and row_vals[-1] is None:
                        row_vals.pop()
                    if not any(v is not None for v in row_vals):
                        continue
                    writer.writerow([str(v) if v is not None else "" for v in row_vals])
            print(f"    -> Saved CSV: {csv_path.name}")
        except Exception as e:
            print(f"    -> Failed to write CSV for sheet {sheet_name}: {e}")
            success = False
            
    return success


def convert_pdf_to_md_basic(pdf_path, output_md_path, images_dir):
    if fitz is None:
        return False, "fitz (PyMuPDF) is not installed."
    
    long_pdf = make_long_path(pdf_path)
    long_md = make_long_path(output_md_path)
    long_images = make_long_path(images_dir)
    
    try:
        doc = fitz.open(long_pdf)
    except Exception as e:
        return False, f"Failed to open PDF: {e}"
        
    text_blocks = []
    image_count = 0
    os.makedirs(long_images, exist_ok=True)
    
    for page_num in range(len(doc)):
        page = doc[page_num]
        text_blocks.append(f"## Page {page_num + 1}\n")
        text_blocks.append(page.get_text("text"))
        
        try:
            image_list = page.get_images(full=True)
            for img_idx, img in enumerate(image_list):
                xref = img[0]
                base_image = doc.extract_image(xref)
                image_bytes = base_image["image"]
                image_ext = base_image["ext"]
                image_name = f"page-{page_num+1}-img-{img_idx+1}.{image_ext}"
                image_path = os.path.join(long_images, image_name)
                
                with open(image_path, "wb") as f_img:
                    f_img.write(image_bytes)
                
                img_dir_name = os.path.basename(images_dir)
                text_blocks.append(f"\n![Ảnh {image_count+1}]({img_dir_name}/{image_name})\n")
                image_count += 1
        except Exception:
            pass
            
    md_content = "\n".join(text_blocks)
    with open(long_md, "w", encoding="utf-8") as f:
        f.write(md_content)
        
    if os.path.exists(long_images) and not os.listdir(long_images):
        os.rmdir(long_images)
        
    return True, f"Success (fitz basic, {image_count} images)"


def convert_pdf_to_md(pdf_path, output_md_path, images_dir, workspace_dir):
    if pymupdf4llm is None or fitz is None:
        return convert_pdf_to_md_basic(pdf_path, output_md_path, images_dir)

    temp_pdf_name = f"_temp_{hashlib.md5(str(pdf_path).encode('utf-8')).hexdigest()[:8]}_convert_.pdf"
    temp_md_name = temp_pdf_name.replace(".pdf", ".md")
    temp_images_name = temp_pdf_name.replace(".pdf", "_images")
    
    temp_pdf_path = workspace_dir / temp_pdf_name
    temp_md_path = workspace_dir / temp_md_name
    temp_images_dir = workspace_dir / temp_images_name
    
    try:
        for path in [temp_pdf_path, temp_md_path]:
            long_p = make_long_path(path)
            if os.path.exists(long_p):
                os.remove(long_p)
        long_t_img = make_long_path(temp_images_dir)
        if os.path.exists(long_t_img):
            shutil.rmtree(long_t_img)
            
        shutil.copy2(make_long_path(pdf_path), make_long_path(temp_pdf_path))
        
        os.makedirs(make_long_path(temp_images_dir), exist_ok=True)
        md_content = pymupdf4llm.to_markdown(
            doc=str(temp_pdf_path),
            write_images=True,
            image_path=str(temp_images_dir),
            image_format="png"
        )
        
        with open(make_long_path(temp_md_path), "w", encoding="utf-8") as f:
            f.write(md_content)
            
        shutil.move(make_long_path(temp_md_path), make_long_path(output_md_path))
        
        long_dest_images = make_long_path(images_dir)
        if os.path.exists(long_t_img) and os.listdir(long_t_img):
            if os.path.exists(long_dest_images):
                shutil.rmtree(long_dest_images)
            shutil.move(long_t_img, long_dest_images)
        elif os.path.exists(long_t_img):
            shutil.rmtree(long_t_img)
            
        return True, "Success (pymupdf4llm via temp copy)"
    except Exception as e:
        print(f"  -> pymupdf4llm failed: {e}. Falling back to basic fitz...")
        return convert_pdf_to_md_basic(pdf_path, output_md_path, images_dir)
    finally:
        try:
            long_temp_pdf = make_long_path(temp_pdf_path)
            if os.path.exists(long_temp_pdf):
                os.remove(long_temp_pdf)
            long_temp_md = make_long_path(temp_md_path)
            if os.path.exists(long_temp_md):
                os.remove(long_temp_md)
            long_t_img = make_long_path(temp_images_dir)
            if os.path.exists(long_t_img):
                shutil.rmtree(long_t_img)
        except Exception:
            pass


def main():
    workspace_dir = Path("d:\\1 Master's Ana Chem\\1 Master's thesis\\Carbon Aerogel")
    target_dir = workspace_dir / "2 Ref lib"
    
    print(f"=== Scanning 2 Ref lib recursive conversions ===")
    
    docx_converted = 0
    docx_skipped = 0
    xlsx_converted = 0
    xlsx_skipped = 0
    pdf_converted = 0
    pdf_skipped = 0
    
    for root, dirs, files in os.walk(make_long_path(target_dir)):
        norm_root = Path(root[4:] if root.startswith('\\\\?\\') else root)
        
        for file in files:
            file_path = norm_root / file
            
            # 1. Process DOCX
            if file.lower().endswith(".docx") and not file.startswith("~$"):
                md_path = file_path.with_suffix(".md")
                if md_path.exists() and md_path.stat().st_size > 10:
                    print(f"DOCX already has MD: {file} -> skipping")
                    docx_skipped += 1
                else:
                    convert_docx_to_md(file_path, md_path)
                    docx_converted += 1
                    
            # 2. Process XLSX
            elif file.lower().endswith(".xlsx") and not file.startswith("~$"):
                # Check if there is already a CSV with the same stem or stem_sheet.csv
                csv_base = file_path.with_suffix(".csv")
                csv_exists = False
                
                # Check for main file csv or any sheet-specific csv
                parent_files = os.listdir(make_long_path(file_path.parent))
                for pf in parent_files:
                    if pf.lower().endswith(".csv") and pf.lower().startswith(file_path.stem.lower()):
                        csv_exists = True
                        break
                        
                if csv_exists:
                    print(f"XLSX already has CSV(s): {file} -> skipping")
                    xlsx_skipped += 1
                else:
                    convert_xlsx_to_csv(file_path, csv_base)
                    xlsx_converted += 1
                    
            # 3. Process PDF
            elif file.lower().endswith(".pdf") and not file.startswith("_temp_"):
                md_path = file_path.with_suffix(".md")
                if md_path.exists() and md_path.stat().st_size > 10:
                    print(f"PDF already has MD: {file} -> skipping")
                    pdf_skipped += 1
                else:
                    print(f"Converting PDF: {file}...")
                    images_dir = file_path.parent / f"{file_path.stem}_images"
                    success, msg = convert_pdf_to_md(file_path, md_path, images_dir, workspace_dir)
                    print(f"  -> {msg}")
                    pdf_converted += 1
                    
    print("\n=== CONVERSION PROCESS COMPLETED ===")
    print(f"DOCX: {docx_converted} converted, {docx_skipped} skipped")
    print(f"XLSX: {xlsx_converted} converted, {xlsx_skipped} skipped")
    print(f"PDF: {pdf_converted} converted, {pdf_skipped} skipped")


if __name__ == "__main__":
    main()
