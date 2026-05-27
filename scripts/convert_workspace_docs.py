# -*- coding: utf-8 -*-
"""
High-fidelity workspace document converter.
Recursively converts all .docx files to .md and all .xlsx files to .csv
in 1 C.A_ syn workflow and Quy trình AI folders.
"""

import os
import sys
import csv
from pathlib import Path

# Set console encoding to UTF-8 on Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

def make_long_path(path):
    abs_path = os.path.abspath(path)
    if os.name == 'nt' and not abs_path.startswith('\\\\?\\'):
        abs_path = abs_path.replace('/', '\\')
        return '\\\\?\\' + abs_path
    return abs_path

# Try importing docx and openpyxl
try:
    from docx import Document
except ImportError:
    print("Error: 'python-docx' is not installed in the current environment.")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("Error: 'openpyxl' is not installed in the current environment.")
    sys.exit(1)

def convert_docx_to_md(docx_path, md_path):
    print(f"Converting DOCX: {docx_path.name} -> {md_path.name}")
    try:
        doc = Document(make_long_path(docx_path))
    except Exception as e:
        print(f"Failed to read DOCX {docx_path.name}: {e}")
        return False

    md_content = []
    
    # We want to traverse paragraphs and tables in order of their appearance
    # Using python-docx body elements
    body_elements = doc.element.body
    tables_iterator = iter(doc.tables)
    paragraphs_iterator = iter(doc.paragraphs)
    
    # Build list of child elements
    for child in body_elements:
        if child.tag.endswith('p'):
            try:
                p = next(paragraphs_iterator)
                text_runs = []
                for run in p.runs:
                    text = run.text
                    if not text:
                        continue
                    if run.bold:
                        text = f"**{text}**"
                    if run.italic:
                        text = f"*{text}*"
                    text_runs.append(text)
                
                paragraph_text = "".join(text_runs).strip()
                if not paragraph_text:
                    md_content.append("")
                    continue
                
                # Check style for headings or lists
                style_name = p.style.name.lower() if p.style and p.style.name else ""
                if "heading 1" in style_name:
                    md_content.append(f"# {paragraph_text}\n")
                elif "heading 2" in style_name:
                    md_content.append(f"## {paragraph_text}\n")
                elif "heading 3" in style_name:
                    md_content.append(f"### {paragraph_text}\n")
                elif "heading 4" in style_name:
                    md_content.append(f"#### {paragraph_text}\n")
                elif "list bullet" in style_name or paragraph_text.startswith(("- ", "* ", "• ")):
                    # Strip existing bullet character if any
                    clean_text = paragraph_text
                    for prefix in ["- ", "* ", "• "]:
                        if clean_text.startswith(prefix):
                            clean_text = clean_text[len(prefix):]
                            break
                    md_content.append(f"- {clean_text}")
                else:
                    md_content.append(paragraph_text)
            except StopIteration:
                pass
        elif child.tag.endswith('tbl'):
            try:
                table = next(tables_iterator)
                table_lines = []
                # Header row
                if len(table.rows) > 0:
                    headers = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
                    table_lines.append("| " + " | ".join(headers) + " |")
                    # Divider
                    table_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                    
                    # Data rows
                    for row in table.rows[1:]:
                        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                        table_lines.append("| " + " | ".join(cells) + " |")
                md_content.append("\n" + "\n".join(table_lines) + "\n")
            except StopIteration:
                pass

    try:
        with open(make_long_path(md_path), "w", encoding="utf-8") as f:
            f.write("\n\n".join(md_content))
        print(f"Successfully saved MD: {md_path.name}")
        return True
    except Exception as e:
        print(f"Failed to write MD {md_path.name}: {e}")
        return False

def convert_xlsx_to_csv(xlsx_path, csv_base_path):
    print(f"Converting XLSX: {xlsx_path.name}")
    try:
        wb = openpyxl.load_workbook(make_long_path(xlsx_path), data_only=True)
    except Exception as e:
        print(f"Failed to read XLSX {xlsx_path.name}: {e}")
        return False

    success = True
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        # Skip empty sheets
        if sheet.max_row == 1 and sheet.max_column == 1 and sheet.cell(1, 1).value is None:
            continue
            
        csv_path = csv_base_path
        if len(wb.sheetnames) > 1:
            # Append sheet name if workbook has multiple sheets
            csv_path = csv_base_path.with_name(f"{csv_base_path.stem}_{sheet_name}.csv")
        else:
            csv_path = csv_base_path.with_name(f"{csv_base_path.stem}.csv")

        print(f"Saving sheet '{sheet_name}' -> {csv_path.name}")
        
        try:
            with open(make_long_path(csv_path), "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                for r in range(1, sheet.max_row + 1):
                    row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
                    # Strip trailing None values to make the CSV clean
                    while row_vals and row_vals[-1] is None:
                        row_vals.pop()
                    # Skip completely empty rows
                    if not any(v is not None for v in row_vals):
                        continue
                    writer.writerow([str(v) if v is not None else "" for v in row_vals])
            print(f"Successfully saved CSV: {csv_path.name}")
        except Exception as e:
            print(f"Failed to write CSV for sheet {sheet_name}: {e}")
            success = False
            
    return success

def main():
    workspace_dir = Path(os.path.dirname(os.path.abspath(__file__))).parent
    
    folders_to_scan = [
        workspace_dir / "1 C.A_ syn workflow",
        workspace_dir / "Quy trình AI"
    ]
    
    for folder in folders_to_scan:
        if not folder.exists():
            print(f"Directory {folder} does not exist, skipping.")
            continue
            
        print(f"\nScanning directory: {folder}")
        for root, dirs, files in os.walk(make_long_path(folder)):
            # Normalize root path to avoid long path prefix in prints
            norm_root = Path(root[4:] if root.startswith('\\\\?\\') else root)
            
            for file in files:
                file_path = norm_root / file
                
                # Process DOCX
                if file.endswith(".docx") and not file.startswith("~$"):
                    md_path = file_path.with_suffix(".md")
                    # Check if MD already exists
                    if md_path.exists():
                        print(f"MD version already exists for {file}, skipping.")
                    else:
                        convert_docx_to_md(file_path, md_path)
                
                # Process XLSX
                elif file.endswith(".xlsx") and not file.startswith("~$"):
                    csv_base_path = file_path.with_suffix(".csv")
                    # Check if CSV already exists (either main name or main_sheetname.csv)
                    # To be safe, we will perform the conversion to ensure all sheets are exported
                    convert_xlsx_to_csv(file_path, csv_base_path)

if __name__ == "__main__":
    main()
