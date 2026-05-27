import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="444444"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, text, title="GHI CHÚ / THÔNG TIN QUAN TRỌNG", color_hex="1B365D", bg_hex="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only (accent border)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:left w:val="single" w:sz="36" w:space="0" w:color="{color_hex}"/>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"★ {title}\n")
    run_title.font.name = 'Arial'
    run_title.font.bold = True
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = RGBColor.from_string(color_hex)
    
    run_text = p.add_run(text.strip())
    run_text.font.name = 'Arial'
    run_text.font.italic = True
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Deep Navy
        
        # Add a bottom border line to Heading 1
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="1B365D"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
        
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x4A, 0x77, 0x7A) # Muted Teal
        
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33) # Dark Charcoal
        
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.bold = True
        run.font.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

def add_chemical_formula_runs(p, formula_text):
    # This parses a LaTeX chemical formula like NH_4OH, Pb^{2+}, Fe-N_4, FeCl_3\cdot6H_2O, etc.
    # Replace LaTeX symbol commands with actual unicode symbols
    text = formula_text.replace(r'\cdot', '·')
    text = text.replace(r'\approx', '≈')
    text = text.replace(r'\le', '≤')
    text = text.replace(r'\ge', '≥')
    text = text.replace(r'\pm', '±')
    text = text.replace(r'\varnothing', 'Ø')
    text = text.replace(r'^\circ', '°')
    text = text.replace(r'\circ', '°')
    text = text.replace(r'\alpha', 'α')
    text = text.replace(r'\theta', 'θ')
    text = text.replace(r'\Delta', 'Δ')
    text = text.replace(r'\rightarrow', '→')
    text = text.replace(r'\%', '%')
    text = text.replace(r'\text{ }', ' ')
    
    # Strip standard \text{...} from text by regex
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    
    # Determine if it's a math variable (italic) or a chemical species (upright)
    is_variable = False
    raw_stripped = text.strip()
    if len(raw_stripped) == 1 and raw_stripped.isalpha():
        is_variable = True
    elif re.match(r'^(E|t|R|S|V|D|C|I|E_p|Delta E_p|R_ct|S_BET|V_pore|D_pore|C_dl|I_pa|I_pc|I_c|E_dep|t_dep|E_dc|quiet|Quiet|P_0|P/P_0|pH|A|n|T)$', raw_stripped):
        is_variable = True
    elif any(prefix in raw_stripped for prefix in ['E_', 't_', 'R_', 'S_', 'V_', 'D_', 'C_', 'I_', 'E{', 't{']):
        is_variable = True
        
    i = 0
    while i < len(text):
        if text[i] == '_':
            i += 1
            if i < len(text) and text[i] == '{':
                end = text.find('}', i)
                if end != -1:
                    sub_text = text[i+1:end]
                    run = p.add_run(sub_text)
                    run.font.subscript = True
                    run.font.name = 'Georgia'
                    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                    i = end + 1
                else:
                    run = p.add_run('_' + text[i:])
                    run.font.name = 'Georgia'
                    break
            else:
                if i < len(text):
                    sub_text = text[i]
                    run = p.add_run(sub_text)
                    run.font.subscript = True
                    run.font.name = 'Georgia'
                    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                    i += 1
                else:
                    run = p.add_run('_')
                    run.font.name = 'Georgia'
        elif text[i] == '^':
            i += 1
            if i < len(text) and text[i] == '{':
                end = text.find('}', i)
                if end != -1:
                    super_text = text[i+1:end]
                    run = p.add_run(super_text)
                    run.font.superscript = True
                    run.font.name = 'Georgia'
                    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                    i = end + 1
                else:
                    run = p.add_run('^' + text[i:])
                    run.font.name = 'Georgia'
                    break
            else:
                if i < len(text):
                    super_text = text[i]
                    run = p.add_run(super_text)
                    run.font.superscript = True
                    run.font.name = 'Georgia'
                    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                    i += 1
                else:
                    run = p.add_run('^')
                    run.font.name = 'Georgia'
        else:
            # Regular text sequence
            start = i
            while i < len(text) and text[i] not in ('_', '^'):
                i += 1
            reg_text = text[start:i]
            run = p.add_run(reg_text)
            run.font.name = 'Georgia'
            run.font.italic = is_variable
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

def format_rich_text(p, text):
    # Regex to extract bold (**), inline code (`), formulas ($), and regular text
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\$.*?\$)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            t_text = token[2:-2]
            run = p.add_run(t_text)
            run.font.bold = True
            run.font.name = 'Arial'
        elif token.startswith('`') and token.endswith('`'):
            t_text = token[1:-1]
            run = p.add_run(t_text)
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xA3, 0x15, 0x15)
        elif token.startswith('$') and token.endswith('$'):
            t_text = token[1:-1]
            add_chemical_formula_runs(p, t_text)
        else:
            run = p.add_run(token)
            run.font.name = 'Arial'
    
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15

def parse_markdown_to_docx(doc, md_filepath):
    with open(md_filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_content = []
    
    in_table = False
    table_headers = []
    table_rows = []
    
    in_callout = False
    callout_type = "NOTE"
    callout_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # 1. Code Blocks
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                
                pPr = p._p.get_or_add_pPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
                pPr.append(shd)
                
                borders = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="8" w:color="CCCCCC"/></w:pBdr>')
                pPr.append(borders)
                
                code_text = "".join(code_content)
                run = p.add_run(code_text.strip())
                run.font.name = 'Courier New'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_content = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_content.append(line)
            i += 1
            continue
            
        # 2. Callouts
        if stripped.startswith('>'):
            match = re.match(r'^>\s*\[!(NOTE|IMPORTANT|WARNING|CAUTION|TIP)\](.*)', stripped, re.IGNORECASE)
            if match:
                in_callout = True
                callout_type = match.group(1).upper()
                extra_text = match.group(2)
                if extra_text.strip():
                    callout_lines.append(extra_text)
            else:
                if in_callout:
                    val = stripped[1:].strip()
                    callout_lines.append(val)
                else:
                    val = stripped[1:].strip()
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.4)
                    pPr = p._p.get_or_add_pPr()
                    borders = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="CCCCCC"/></w:pBdr>')
                    pPr.append(borders)
                    run = p.add_run(val)
                    run.font.italic = True
                    run.font.name = 'Arial'
                    run.font.size = Pt(10)
            i += 1
            continue
        else:
            if in_callout:
                in_callout = False
                c_text = " ".join(callout_lines)
                
                color_map = {
                    "IMPORTANT": ("8B0000", "LƯU Ý QUAN TRỌNG (CRITICAL)", "FDF2F2"),
                    "WARNING": ("B25900", "CẢNH BÁO (WARNING / HAZARD)", "FFF9F2"),
                    "CAUTION": ("B25900", "THẬN TRỌNG (CAUTION)", "FFF9F2"),
                    "TIP": ("006400", "MẸO KỸ THUẬT (EXPERT TIP)", "F2FDF2"),
                    "NOTE": ("1B365D", "THAM CHIẾU (NOTE / REFERENCE)", "F0F4F8")
                }
                color_hex, title_str, bg_hex = color_map.get(callout_type, ("1B365D", "GHI CHÚ / THÔNG TIN", "F0F4F8"))
                
                add_callout_box(doc, c_text, title=title_str, color_hex=color_hex, bg_hex=bg_hex)
                callout_lines = []
                
        # 3. Headings
        if stripped.startswith('#'):
            heading_match = re.match(r'^(#+)\s*(.*)', stripped)
            if heading_match:
                level = len(heading_match.group(1))
                h_text = heading_match.group(2).strip()
                if level == 1 and h_text.endswith('.md'):
                    h_text = h_text.split('—')[1].strip() if '—' in h_text else h_text
                add_styled_heading(doc, h_text, level)
            i += 1
            continue
            
        # 4. Tables
        if stripped.startswith('|'):
            in_table = True
            row_cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if all(re.match(r'^[-:\s]+$', c) for c in row_cells):
                i += 1
                continue
            
            if len(table_headers) == 0:
                table_headers = row_cells
            else:
                table_rows.append(row_cells)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                col_count = len(table_headers)
                if col_count > 0:
                    tbl = doc.add_table(rows=len(table_rows) + 1, cols=col_count)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.autofit = True
                    set_table_borders(tbl)
                    
                    hdr_row = tbl.rows[0]
                    trPr = hdr_row._tr.get_or_add_trPr()
                    trPr.append(OxmlElement('w:tblHeader'))
                    trPr.append(OxmlElement('w:cantSplit'))
                    
                    for col_idx, text in enumerate(table_headers):
                        cell = hdr_row.cells[col_idx]
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_background(cell, "1B365D")
                        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                        
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        p.paragraph_format.space_after = Pt(0)
                        p.paragraph_format.space_before = Pt(0)
                        run = p.add_run(text)
                        run.font.bold = True
                        run.font.name = 'Arial'
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        
                    for row_idx, row_data in enumerate(table_rows):
                        row = tbl.rows[row_idx + 1]
                        trPr = row._tr.get_or_add_trPr()
                        trPr.append(OxmlElement('w:cantSplit'))
                        
                        bg_color = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
                        
                        for col_idx, text in enumerate(row_data):
                            if col_idx < len(row.cells):
                                cell = row.cells[col_idx]
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                set_cell_background(cell, bg_color)
                                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                                
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_after = Pt(0)
                                p.paragraph_format.space_before = Pt(0)
                                format_rich_text(p, text)
                                p.runs[0].font.size = Pt(9)
                                
                table_headers = []
                table_rows = []
                
        # 5. Lists
        bullet_match = re.match(r'^([\*\-\+])\s+(.*)', stripped)
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        
        if bullet_match:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            content = bullet_match.group(2).strip()
            format_rich_text(p, content)
            i += 1
            continue
            
        elif num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            content = num_match.group(2).strip()
            format_rich_text(p, content)
            i += 1
            continue
            
        # 6. Regular paragraphs
        if stripped:
            p = doc.add_paragraph()
            format_rich_text(p, stripped)
            
        i += 1

def main():
    doc = Document()
    
    # Document Margins (Standard 1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # --- Title Page (Trang Bìa) ---
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(72)
    p_top.paragraph_format.space_after = Pt(12)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("QUY TRÌNH TIÊU CHUẨN ĐỐI CHỨNG HỢP NHẤT\nCHẾ TẠO CARBON AEROGEL ĐIỆN HÓA CẢM BIẾN")
    run_title.font.name = 'Arial'
    run_title.font.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    p_title.paragraph_format.space_after = Pt(18)
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_subtitle.add_run(
        "Tài liệu Tích hợp, Hợp nhất, Đối chứng định lượng % lý tính và Biện luận Cơ chế\n"
        "Tách chiết Cellulose I xơ dừa Bến Tre ──> Sol-gel chuyển pha Cellulose II & III ──>\n"
        "Tẩm Sắt Co-doped ──> Sấy chân không đông khô ──> Carbon hóa bóc tách mạng sp² ──>\n"
        "Hòa tan cặn axit leaching ──> Cấu hình phối trí đơn nguyên tử Sắt-Nitơ (Fe-N₄)."
    )
    run_sub.font.name = 'Arial'
    run_sub.font.italic = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p_subtitle.paragraph_format.space_after = Pt(100)
    
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_info = p_info.add_run(
        "TÀI LIỆU KHUNG LUẬN VĂN THẠC SĨ HÓA HỌC\n"
        "Học viên: LÊ HOÀNG NHẬT\n"
        "Dựa trên các nghiên cứu cột mốc:\n"
        "Nguyễn Trần Xuân Phương (2024) | Fauziyah et al. (2020) | Wu et al. (2024)\n"
        "Thành phố Hồ Chí Minh, Việt Nam"
    )
    run_info.font.name = 'Arial'
    run_info.font.bold = True
    run_info.font.size = Pt(10.5)
    run_info.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    doc.add_page_break()
    
    # --- Parse the Final Consolidated SOP md file ---
    filepath = r"D:\1 Master's Ana Chem\1 Master's thesis\Carbon Aerogel\1 SOPs\Final_SOPs_CA.md"
    if os.path.exists(filepath):
        print(f"Parsing consolidated SOP file {filepath}...")
        parse_markdown_to_docx(doc, filepath)
        
    dest_path1 = r"D:\1 Master's Ana Chem\1 Master's thesis\Carbon Aerogel\1 SOPs\Final_SOPs_CA.docx"
    dest_path2 = r"D:\1 Master's Ana Chem\1 Master's thesis\Carbon Aerogel\Final_SOPs_CA.docx"
    
    doc.save(dest_path1)
    doc.save(dest_path2)
    print(f"Successfully compiled and saved integrated document to:\n1. {dest_path1}\n2. {dest_path2}")

if __name__ == "__main__":
    main()
