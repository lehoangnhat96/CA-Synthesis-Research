import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def check_renumbered():
    doc = docx.Document(r"DCLV_CA_02.07 Fe-N CA_Renumbered.docx")
    
    print("--- Checking Table 1.5 ---")
    for table in doc.tables:
        for row in table.rows:
            if 'Bismuth-Chitosan Nanocomposite' in row.cells[0].text or 'rGO-Chitosan/PLL Nanocomposite' in row.cells[0].text:
                print("Row found:", [cell.text.strip() for cell in row.cells])
                
    print("\n--- Checking Bibliography end ---")
    bib_entries = []
    in_bib = False
    for p in doc.paragraphs:
        if "TÀI LIỆU THAM KHẢO" in p.text.upper():
            in_bib = True
        if in_bib and p.text.strip().startswith('['):
            bib_entries.append(p.text.strip())
            
    if bib_entries:
        for entry in bib_entries[-5:]:
            print(entry)
            
if __name__ == "__main__":
    check_renumbered()
