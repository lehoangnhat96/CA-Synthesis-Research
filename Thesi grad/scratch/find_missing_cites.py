import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def find_missing():
    # 1. DOCX
    doc = docx.Document(r"DCLV_CA_02.07 Fe-N CA.docx")
    docx_cites = set()
    
    def process_docx(text):
        matches = re.finditer(r'\[([\d,\s\-]+)\]', text)
        for m in matches:
            for part in m.group(1).split(','):
                part = part.strip()
                if '-' in part:
                    try:
                        start, end = map(int, part.split('-'))
                        for num in range(start, end + 1):
                            docx_cites.add(num)
                    except ValueError: pass
                elif part.isdigit():
                    docx_cites.add(int(part))

    for p in doc.paragraphs:
        text = p.text.strip()
        if not re.match(r'^\[(\d+)\]\s+(.*)', text):
            process_docx(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_docx(cell.text.strip())
                
    # 2. MD
    with open(r"DCLV_CA_01.07 Fe-N CA.md", 'r', encoding='utf-8') as f:
        md_lines = f.readlines()
        
    md_cites = set()
    for text in md_lines:
        text = text.strip()
        if re.match(r'^\\\[(\d+)\\\]\s+(.*)', text):
            continue
        matches = re.finditer(r'\\?\[([\d,\s\-]+)\\?\](?!\()', text)
        for m in matches:
            for part in m.group(1).split(','):
                part = part.strip()
                if '-' in part:
                    try:
                        start, end = map(int, part.split('-'))
                        for num in range(start, end + 1):
                            md_cites.add(num)
                    except ValueError: pass
                elif part.isdigit():
                    md_cites.add(int(part))
                    
    missing_in_md = docx_cites - md_cites
    print("Citations in DOCX but not found by regex in MD:", sorted(list(missing_in_md)))

if __name__ == "__main__":
    find_missing()
