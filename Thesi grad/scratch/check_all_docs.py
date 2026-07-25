import docx
import re
import sys

def iter_block_items(parent):
    from docx.document import Document
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.table import _Cell, Table
    from docx.text.paragraph import Paragraph
    
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Something's wrong")
        
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            table = Table(child, parent)
            yield table

def parse_citations(doc_path):
    doc = docx.Document(doc_path)
    citations_in_order = []
    seen = set()
    in_refs = False
    
    cit_pattern = re.compile(r'\[([0-9\s,\-]+)\]')
    
    def process_text(text):
        nonlocal in_refs
        if 'TÀI LIỆU THAM KHẢO' in text:
            in_refs = True
            
        if not in_refs:
            matches = cit_pattern.findall(text)
            for match in matches:
                parts = match.split(',')
                for part in parts:
                    part = part.strip()
                    if not part: continue
                    if '-' in part:
                        try:
                            start, end = map(int, part.split('-'))
                            for i in range(start, end + 1):
                                if i not in seen:
                                    seen.add(i)
                                    citations_in_order.append(i)
                        except ValueError:
                            pass
                    else:
                        try:
                            num = int(part)
                            if num not in seen:
                                seen.add(num)
                                citations_in_order.append(num)
                        except ValueError:
                            pass

    for block in iter_block_items(doc):
        if in_refs:
            break
        if block.__class__.__name__ == 'Paragraph':
            process_text(block.text)
        elif block.__class__.__name__ == 'Table':
            for row in block.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        process_text(p.text)
                        
    refs = []
    for p in doc.paragraphs:
        if 'TÀI LIỆU THAM KHẢO' in p.text:
            in_refs = True
            continue
        if in_refs:
            m = re.match(r'^\[(\d+)\]', p.text.strip())
            if m:
                refs.append(int(m.group(1)))
                
    print(f"File: {doc_path}")
    print(f"Total unique citations in text: {len(seen)}")
    print(f"Order of first appearance: {citations_in_order}")
    
    mismatches = 0
    for i, val in enumerate(citations_in_order):
        if i + 1 != val:
            mismatches += 1
            
    print(f"Number of out-of-order citations: {mismatches}")
    print("-" * 40)

if __name__ == '__main__':
    parse_citations('DCLV_CA_29.06 Fe-N CA.docx')
    parse_citations('DCLV_CA_29.06 Fe-N CA_Reordered.docx')
    parse_citations('DCLV_CA_29.06 Fe-N CA_Reordered_V2.docx')
