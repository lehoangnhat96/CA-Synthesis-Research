import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def replace_text_in_paragraph(p, old_text, new_text):
    if old_text in p.text:
        # Simple text replacement loses formatting, but for these specific blocks it's acceptable
        # A more complex replacement preserves runs
        inline = p.runs
        for i in range(len(inline)):
            if old_text in inline[i].text:
                inline[i].text = inline[i].text.replace(old_text, new_text)
                return True
        # If split across runs, we do a full replacement of the paragraph text
        p.text = p.text.replace(old_text, new_text)
        return True
    return False

def main():
    doc_path = r"d:\1 Master's Ana Chem\1 Master's thesis\Carbon Aerogel\Thesi grad\De_cuong_Luan_van_Thac_si_Le_Hoang_Nhat.docx.docx"
    doc = docx.Document(doc_path)
    
    # 1. Fix Table 1.2
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if 'Cellulose III' in cell.text and 'Cellulose II' in cell.text:
                    if 'Pha tinh thể ẩm' in cell.text:
                        for p in cell.paragraphs:
                            if 'Cellulose III' in p.text:
                                p.text = p.text.replace('Cellulose III', 'Cellulose II / Na-cellulose IV (giả thuyết — [ANALOGICAL-MECHANISTIC], chưa xác nhận XRD trên hệ NH4OH)')

    # 2. Fix 2.4.2.2 paragraph
    target_phrase = "lộ trình phức hợp amoniac-cellulose"
    new_para_text = """Dưới lực ép cơ học của tinh thể đá đang phát triển và cự ly tiếp xúc phân tử cực ngắn (< 0.3 nm), lớp kén bảo vệ của urea bị phá vỡ, các nhóm hydroxyl tự do trên các chuỗi cellulose lân cận tự kết nối chéo vật lý trực tiếp với nhau.

[ANALOGICAL-MECHANISTIC] Đối chiếu với dữ liệu synchrotron XRD của Isobe et al. trên hệ LiOH/urea đông lạnh [34], cơ chế kết tinh lại dạng này nhiều khả năng dẫn đến Cellulose II hoặc pha trung gian Na-cellulose IV, chứ không phải Cellulose III — vì Cellulose III về nguyên tắc chỉ hình thành từ tác nhân trương nở khan (ammonia lỏng khan, ethylenediamine...), trong khi hệ của đề tài dùng NH4OH 25% ở dạng dung dịch nước.

[NO_DIRECT_EVIDENCE] Isobe et al. [34] thực nghiệm trên hệ LiOH/urea, không phải NH4OH/urea; việc ngoại suy sang hệ của đề tài chưa có bằng chứng trực tiếp và cần được xác nhận bằng XRD trên mẫu pre-pyrolysis (trước GĐ5, vì nhiệt phân phá hủy cấu trúc tinh thể cellulose) trước khi kết luận pha tinh thể cuối cùng.

Khung gel 3D tổ ong được định hình trong quá trình này, giam giữ các tinh thể đá làm khuôn lỗ xốp lớn (macropores)."""

    for p in doc.paragraphs:
        if target_phrase in p.text:
            p.text = new_para_text
            break

    # 3. Fix Figure 2.5 Caption
    for p in doc.paragraphs:
        if 'Hình 2.5' in p.text and 'Cellulose III' in p.text:
            replace_text_in_paragraph(p, 'Cellulose III', 'Cellulose II (hoặc Na-cellulose IV)')
            
    # 4. Fix 2.4.4.2 
    for p in doc.paragraphs:
        if 'vách xốp cellulose III mỏng' in p.text:
            replace_text_in_paragraph(p, 'vách xốp cellulose III mỏng', 'vách xốp cellulose mỏng')
            
    # 5. Fix 2.4.2.2 Heading if it says Cellulose III
    for p in doc.paragraphs:
        if 'Sự tái tổ chức Cellulose III' in p.text:
            replace_text_in_paragraph(p, 'Sự tái tổ chức Cellulose III', 'Sự tái tổ chức Cellulose II / Na-cellulose IV')

    out_path = r"d:\1 Master's Ana Chem\1 Master's thesis\Carbon Aerogel\Thesi grad\De_cuong_Luan_van_Thac_si_Le_Hoang_Nhat_Fixed_Cellulose.docx"
    doc.save(out_path)
    print("Fixed document saved to:", out_path)

if __name__ == '__main__':
    main()
