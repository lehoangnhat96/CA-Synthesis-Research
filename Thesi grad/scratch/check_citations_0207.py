import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def extract_citations():
    doc_path = r"DCLV_CA_02.07 Fe-N CA_Renumbered.docx"
    doc = docx.Document(doc_path)
    
    bibliography = {}
    seen_in_text = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        
        # Match [N] at start of paragraph (Bibliography entry)
        bib_match = re.match(r'^\[(\d+)\]\s+(.*)', text)
        if bib_match:
            num = int(bib_match.group(1))
            content = bib_match.group(2).strip()
            bibliography[num] = content
            continue
            
        # Match in-text citations [N]
        matches = re.finditer(r'\[([\d,\s\-]+)\]', text)
        for m in matches:
            inner = m.group(1)
            parts = inner.split(',')
            for part in parts:
                part = part.strip()
                if '-' in part:
                    try:
                        start, end = map(int, part.split('-'))
                        for num in range(start, end + 1):
                            if num not in seen_in_text:
                                seen_in_text.append(num)
                    except ValueError:
                        pass
                else:
                    if part.isdigit():
                        num = int(part)
                        if num not in seen_in_text:
                            seen_in_text.append(num)
                            
    # Search tables (tables appear in the body)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                matches = re.finditer(r'\[([\d,\s\-]+)\]', text)
                for m in matches:
                    inner = m.group(1)
                    parts = inner.split(',')
                    for part in parts:
                        part = part.strip()
                        if '-' in part:
                            try:
                                start, end = map(int, part.split('-'))
                                for num in range(start, end + 1):
                                    if num not in seen_in_text:
                                        seen_in_text.append(num)
                            except ValueError:
                                pass
                        else:
                            if part.isdigit():
                                num = int(part)
                                if num not in seen_in_text:
                                    seen_in_text.append(num)

    # 3. Analyze Discrepancies
    print("=== ANALYSIS RESULTS ===")
    print(f"Total citations found in text: {len(seen_in_text)}")
    print(f"Total entries in bibliography: {len(bibliography)}")
    
    out_of_order = []
    missing_in_bib = []
    
    expected_next = 1
    for num in seen_in_text:
        if num != expected_next:
            if num > expected_next:
                out_of_order.append((num, expected_next))
        expected_next = max(expected_next, num) + 1
        
        if num not in bibliography:
            missing_in_bib.append(num)
            
    print("\n--- Out of Order Citations ---")
    if out_of_order:
        for actual, expected in out_of_order:
            print(f"Found [{actual}], but expected [{expected}] based on appearance order.")
    else:
        print("All citations appear in strict sequential order!")
        
    print("\n--- Citations missing from Bibliography ---")
    if missing_in_bib:
        for num in missing_in_bib:
            print(f"[{num}] appears in text but has no bibliography entry!")
    else:
        print("All text citations have a corresponding bibliography entry.")

    print("\n--- Bibliography entries never cited in text ---")
    unused = set(bibliography.keys()) - set(seen_in_text)
    if unused:
        for num in sorted(unused):
            print(f"[{num}] is in bibliography but never cited in text.")
    else:
        print("All bibliography entries are cited at least once.")

if __name__ == "__main__":
    extract_citations()
