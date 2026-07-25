import docx
import re

def parse_citations():
    doc_path = 'DCLV_CA_29.06 Fe-N CA_Reordered_V2.docx'
    doc = docx.Document(doc_path)
    
    citations_in_order = []
    seen = set()
    
    in_refs = False
    
    # regex for citations like [1], [1, 2], [1-3], [1, 3-5]
    cit_pattern = re.compile(r'\[([0-9\s,\-]+)\]')
    
    for p in doc.paragraphs:
        if 'TÀI LIỆU THAM KHẢO' in p.text:
            in_refs = True
            break
            
        text = p.text
        matches = cit_pattern.findall(text)
        for match in matches:
            # parse the match string
            parts = match.split(',')
            for part in parts:
                part = part.strip()
                if not part: continue
                if '-' in part:
                    try:
                        start, end = map(int, part.split('-'))
                        for i in range(start, end + 1):
                            if i not in seen:
                                seen.add(i)
                                citations_in_order.append(i)
                    except ValueError:
                        pass
                else:
                    try:
                        num = int(part)
                        if num not in seen:
                            seen.add(num)
                            citations_in_order.append(num)
                    except ValueError:
                        pass
                        
    # Now check the bibliography
    refs = []
    for p in doc.paragraphs:
        if 'TÀI LIỆU THAM KHẢO' in p.text:
            in_refs = True
            continue
        if in_refs:
            m = re.match(r'^\[(\d+)\]', p.text.strip())
            if m:
                refs.append(int(m.group(1)))
                
    with open('citation_check_report.txt', 'w', encoding='utf-8') as f:
        f.write(f"Total unique citations in text: {len(seen)}\n")
        f.write(f"Order of first appearance: {citations_in_order}\n")
        f.write(f"Expected order: {list(range(1, len(seen)+1))}\n")
        
        mismatches = []
        for i, val in enumerate(citations_in_order):
            if i + 1 != val:
                mismatches.append(f"Expected {i+1}, found {val}")
                
        if mismatches:
            f.write("ORDER ISSUES IN TEXT:\n")
            for m in mismatches[:10]:
                f.write("  " + m + "\n")
            if len(mismatches) > 10:
                f.write(f"  ... and {len(mismatches)-10} more.\n")
        else:
            f.write("No order issues in text. Citations appear sequentially.\n")
            
        f.write(f"\nTotal references in bibliography: {len(refs)}\n")
        missing_in_bib = set(citations_in_order) - set(refs)
        missing_in_text = set(refs) - set(citations_in_order)
        
        if missing_in_bib:
            f.write(f"Citations in text but missing in bibliography: {sorted(list(missing_in_bib))}\n")
        if missing_in_text:
            f.write(f"Citations in bibliography but not found in text: {sorted(list(missing_in_text))}\n")

if __name__ == '__main__':
    parse_citations()
