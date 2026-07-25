import shutil
import datetime
import sys
import docx
from copy import deepcopy

sys.stdout.reconfigure(encoding='utf-8')

doc_path = r"DCLV_CA_01.07 Fe-N CA.docx"

timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = doc_path + f".bak_{timestamp}"
try:
    shutil.copy2(doc_path, backup_path)
    print(f"[OK] Backup saved: {backup_path}")
except Exception as e:
    print(f"[FATAL] Cannot create backup: {e}. ABORTED.")
    sys.exit(1)

doc = docx.Document(doc_path)

# 1. Update Bibliography Entry [62]
for i, p in enumerate(doc.paragraphs):
    if '[62] Y. El Hamdouni' in p.text:
        # Assuming it's in the first run as checked
        p.runs[0].text = '[62] Z. Guo et al., "Simultaneous determination of trace Cd(II), Pb(II) and Cu(II) by differential pulse anodic stripping voltammetry using a reduced graphene oxide-chitosan/poly-L-lysine nanocomposite modified glassy carbon electrode," J. Colloid Interface Sci., vol. 490, pp. 11-22, Mar. 2017, doi: 10.1016/j.jcis.2016.11.006.'
        print(f"[OK] Bibliography entry [62] replaced at paragraph {i}.")
        break

# 2. Insert into Table 1.5
found_table = False
for tbl in doc.tables:
    for idx, row in enumerate(tbl.rows):
        if len(row.cells) == 6 and 'Bismuth-Chitosan Nanocomposite' in row.cells[0].text:
            new_tr = deepcopy(row._tr)
            row._tr.addprevious(new_tr)
            new_row = tbl.rows[idx] 
            new_row.cells[0].text = 'rGO-Chitosan/PLL Nanocomposite'
            new_row.cells[1].text = 'Pb²⁺, Cd²⁺, Cu²⁺'
            new_row.cells[2].text = 'Chitosan'
            new_row.cells[3].text = '~0.0001 µM'
            new_row.cells[4].text = '0.002 - 0.38 µM'
            new_row.cells[5].text = '[62]'
            
            if len(new_row.cells[5].paragraphs) > 0:
                new_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                
            print("[OK] Table 1.5 updated with row [62] Z. Guo 2017.")
            found_table = True
            break
    if found_table:
        break

doc.save(doc_path)
print("[OK] Saved successfully to " + doc_path)
