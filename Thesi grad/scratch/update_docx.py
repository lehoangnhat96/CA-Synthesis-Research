import shutil
import datetime
import sys
import docx
from copy import deepcopy

sys.stdout.reconfigure(encoding='utf-8')

# SOURCE FILE: DCLV_CA_01.07 Fe-N CA.docx
doc_path = r"DCLV_CA_01.07 Fe-N CA.docx"

timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = doc_path + f".bak_{timestamp}"
try:
    shutil.copy2(doc_path, backup_path)
    print(f"[OK] Backup saved: {backup_path}")
except Exception as e:
    print(f"[FATAL] Cannot create backup: {e}. ABORTED.")
    sys.exit(1)

doc = docx.Document(doc_path)

# TEXT REPLACEMENTS IN RUNS
# 1. Paragraph 183
p183 = doc.paragraphs[183]
old_str_1 = 'định hình các liên kết Fe-N bền chặt [34, 35]'
new_str_1 = 'định hình cấu trúc liên kết Fe-N (điều này đã được minh chứng thông qua các mô hình phân tử mô phỏng lý thuyết [34] và trên vật liệu thực tế [35])'
if old_str_1 in p183.runs[1].text:
    p183.runs[1].text = p183.runs[1].text.replace(old_str_1, new_str_1)
    print("[OK] Paragraph 183 Run 1 updated.")

old_str_2 = 'Trong đó, một nghiên cứu gần đây về carbon aerogel pha tạp Fe và N đã chứng minh vật liệu này có khả năng cảm biến điện hóa đa phương thức (dual-mode) hiệu quả, đặc biệt là trong việc phát hiện glucose với giới hạn phát hiện (LOD) đạt 0.5 µM [36].'
new_str_2 = 'Đây được xem là bằng chứng trực tiếp mạnh mẽ nhất cho thấy nền tảng vật liệu aerogel pha tạp Fe/N hoàn toàn có khả năng hoạt động như một cảm biến sinh học nhạy bén. Chẳng hạn, một nghiên cứu thiết kế Fe/N-CA phân cấp đã cho thấy khả năng phát hiện đa phương thức đối với phân tử glucose với giới hạn phát hiện (LOD) rất thấp đạt 0.5 µM [36]. Tuy nhiên, việc phát triển hệ cảm biến này từ nguồn sinh khối xơ dừa giá rẻ và định hướng tích hợp màng Chitosan/Nafion để kiểm soát nhiễu sinh học, đặc biệt hướng tới phân tích phức hợp dược phẩm (Paracetamol) và kim loại nặng, vẫn là một khoảng trống nghiên cứu (research gap) đầy triển vọng.'
if old_str_2 in p183.runs[5].text:
    p183.runs[5].text = p183.runs[5].text.replace(old_str_2, new_str_2)
    print("[OK] Paragraph 183 Run 5 updated.")

# 2. Paragraph 303
p303 = doc.paragraphs[303]
old_str_3 = 'thông qua phổ XPS, phổ Raman và phổ tổng trở điện hóa EIS [34, 35]'
new_str_3 = 'thông qua phổ XPS, phổ Raman và phổ tổng trở điện hóa EIS [35]'
if old_str_3 in p303.runs[6].text:
    p303.runs[6].text = p303.runs[6].text.replace(old_str_3, new_str_3)
    print("[OK] Paragraph 303 Run 6 updated.")

# 3. Paragraph 436
p436 = doc.paragraphs[436]
old_str_4 = 'phẩm có dòng oxi hóa phân hủy phức tạp như Paracetamol [55, 61, 65]."'
new_str_4 = 'phẩm có dòng oxi hóa phân hủy phức tạp như Paracetamol [55, 61]. Hơn nữa, việc sử dụng các điện cực biến tính bằng màng Nafion-composite (cụ thể là Nafion/TiO₂-graphene) đã thiết lập một hệ quy chiếu hiệu năng quan trọng cho phép phân tích Paracetamol với độ nhạy cao và LOD đạt mức vết (0.21 µM) [65], tạo cơ sở trực tiếp để so sánh và định chuẩn (benchmark) hiệu năng cảm biến DPV trong đề tài này."'
if old_str_4 in p436.runs[0].text:
    p436.runs[0].text = p436.runs[0].text.replace(old_str_4, new_str_4)
    print("[OK] Paragraph 436 Run 0 updated.")
else:
    # try replacing without the quote
    old_str_4_alt = 'phẩm có dòng oxi hóa phân hủy phức tạp như Paracetamol [55, 61, 65].'
    new_str_4_alt = 'phẩm có dòng oxi hóa phân hủy phức tạp như Paracetamol [55, 61]. Hơn nữa, việc sử dụng các điện cực biến tính bằng màng Nafion-composite (cụ thể là Nafion/TiO₂-graphene) đã thiết lập một hệ quy chiếu hiệu năng quan trọng cho phép phân tích Paracetamol với độ nhạy cao và LOD đạt mức vết (0.21 µM) [65], tạo cơ sở trực tiếp để so sánh và định chuẩn (benchmark) hiệu năng cảm biến DPV trong đề tài này.'
    if old_str_4_alt in p436.runs[0].text:
        p436.runs[0].text = p436.runs[0].text.replace(old_str_4_alt, new_str_4_alt)
        print("[OK] Paragraph 436 Run 0 updated (no quote).")


# TABLE 1.5 INSERTION
found_table = False
for tbl in doc.tables:
    for idx, row in enumerate(tbl.rows):
        if len(row.cells) > 0 and 'Fe/N-CA' in row.cells[0].text:
            new_tr = deepcopy(row._tr)
            row._tr.addprevious(new_tr)
            new_row = tbl.rows[idx] # the newly inserted row shifts into this index
            new_row.cells[0].text = 'Nafion/TiO₂-graphene/GCE'
            new_row.cells[1].text = 'Paracetamol'
            new_row.cells[2].text = 'Nafion'
            new_row.cells[3].text = '0.21 µM'
            new_row.cells[4].text = '1 – 100 µM'
            new_row.cells[5].text = '[65]'
            
            # center align the last cell
            if len(new_row.cells[5].paragraphs) > 0:
                new_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                
            print("[OK] Table 1.5 updated with row [65].")
            found_table = True
            break
    if found_table:
        break

doc.save(doc_path)
print("[OK] Saved successfully to " + doc_path)
