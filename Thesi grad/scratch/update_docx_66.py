import shutil
import datetime
import sys
import docx
from copy import deepcopy

sys.stdout.reconfigure(encoding='utf-8')

# QT-03: Backup
doc_path = r"DCLV_CA_01.07 Fe-N CA.docx"
timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
backup_path = doc_path + f".bak_{timestamp}"
try:
    shutil.copy2(doc_path, backup_path)
    print(f"[OK] Backup saved: {backup_path}")
except Exception as e:
    print(f"[FATAL] Cannot create backup: {e}. ABORTED.")
    sys.exit(1)

doc = docx.Document(doc_path)

# Task 1: Insert into Table 1.5
found_table = False
for tbl in doc.tables:
    for idx, row in enumerate(tbl.rows):
        if len(row.cells) == 6 and 'Bismuth-Chitosan Nanocomposite' in row.cells[0].text:
            new_tr = deepcopy(row._tr)
            row._tr.addprevious(new_tr)
            new_row = tbl.rows[idx] 
            new_row.cells[0].text = 'rGO-Chitosan/PLL Nanocomposite'
            new_row.cells[1].text = 'Pb²⁺, Cd²⁺, Cu²⁺'
            new_row.cells[2].text = 'Chitosan'
            new_row.cells[3].text = '~0.0001 µM'
            new_row.cells[4].text = '0.002 - 0.38 µM'
            new_row.cells[5].text = '[66]'
            
            if len(new_row.cells[5].paragraphs) > 0:
                new_row.cells[5].paragraphs[0].alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
                
            print("[OK] Table 1.5 updated with row [66] Z. Guo 2017.")
            found_table = True
            break
    if found_table:
        break

# Task 2: Update Paragraphs in text
text_insert_1 = " Đặc biệt, việc sử dụng màng Chitosan không chỉ đóng vai trò chất kết dính cơ học mà còn tăng cường tín hiệu hòa tan anode nhờ khả năng tạo phức chelate mạnh mẽ của các nhóm -NH₂ với ion kim loại. Cơ chế hiệp đồng này tương tự như thiết kế đã được Guo và cộng sự chứng minh thành công trên hệ nền RGO-Chitosan [66]."
text_insert_2 = " Để giải quyết rủi ro này, sự kết hợp giữa nền carbon xốp (Fe/N-CA) và Chitosan được kỳ vọng sẽ tách bạch rõ ràng các đỉnh điện thế của Cd, Pb, Zn, ngăn ngừa sự tạo hợp kim liên kim loại (intermetallic compounds), dựa trên thành công của các thiết kế cấu trúc carbon xốp/chitosan tương tự [66]."

for p in doc.paragraphs:
    if "tăng cường quá trình làm giàu (accumulation) các cation kim loại nặng" in p.text and "[60, 61]" in p.text:
        # Append to paragraph
        p.add_run(text_insert_1)
        print("[OK] Paragraph 1 (Chitosan Chelate Mechanism) updated.")
        
    if "như Pb²⁺ và Cd²⁺" in p.text and "Sự trôi thế của điện cực so sánh" in p.text:
        # Insert into middle of paragraph after "(như Pb²⁺ và Cd²⁺)."
        # This is tricky with runs, so we will do a targeted string replacement in text but we lose formatting. 
        # But wait! We can just find the run containing "và Cd²⁺)." and append to it.
        for run in p.runs:
            if "và Cd²⁺)." in run.text:
                run.text = run.text.replace("và Cd²⁺).", "và Cd²⁺)." + text_insert_2)
                print("[OK] Paragraph 2 (Simultaneous Detection) updated.")
                break

# Task 3: Append to Bibliography
# We find [65] and append a new paragraph after it.
# Actually doc.paragraphs doesn't allow insert_paragraph_after easily without _element.addnext, let's just find [65] and use insert_paragraph_before on the following paragraph.
bib_entry = '[66] Z. Guo et al., "Simultaneous determination of trace Cd(II), Pb(II) and Cu(II) by differential pulse anodic stripping voltammetry using a reduced graphene oxide-chitosan/poly-L-lysine nanocomposite modified glassy carbon electrode," J. Colloid Interface Sci., vol. 490, pp. 11–22, Mar. 2017, doi: 10.1016/j.jcis.2016.11.006.'

for i, p in enumerate(doc.paragraphs):
    if "[65] F. Yang" in p.text:
        # The next paragraph is likely empty or the next section.
        next_p = doc.paragraphs[i+1]
        new_p = next_p.insert_paragraph_before(bib_entry)
        # Match formatting
        new_p.style = p.style
        print("[OK] Bibliography appended [66].")
        break

doc.save(doc_path)
print("[OK] Saved successfully to " + doc_path)
