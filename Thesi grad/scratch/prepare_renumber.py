import docx
import re
import sys

sys.stdout.reconfigure(encoding='utf-8')

def plan_renumber():
    doc_path = r"DCLV_CA_02.07 Fe-N CA.docx"
    doc = docx.Document(doc_path)
    
    bibliography = {}
    seen_in_text = []
    
    # 1. Extract Bibliography
    for p in doc.paragraphs:
        text = p.text.strip()
        bib_match = re.match(r'^\[(\d+)\]\s+(.*)', text)
        if bib_match:
            num = int(bib_match.group(1))
            content = bib_match.group(2).strip()
            bibliography[num] = content
            
    # 2. Extract In-Text sequential order
    def process_text_for_cites(text):
        matches = re.finditer(r'\[([\d,\s\-]+)\]', text)
        for m in matches:
            inner = m.group(1)
            parts = inner.split(',')
            for part in parts:
                part = part.strip()
                if '-' in part:
                    try:
                        start, end = map(int, part.split('-'))
                        for num in range(start, end + 1):
                            if num not in seen_in_text:
                                seen_in_text.append(num)
                    except ValueError:
                        pass
                else:
                    if part.isdigit():
                        num = int(part)
                        if num not in seen_in_text:
                            seen_in_text.append(num)

    for p in doc.paragraphs:
        text = p.text.strip()
        if re.match(r'^\[(\d+)\]\s+(.*)', text):
            continue # Skip bibliography definitions
        process_text_for_cites(text)
        
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                process_text_for_cites(cell.text.strip())
                
    # 3. Create mapping Old -> New
    old_to_new = {}
    new_to_old = {}
    
    current_new_id = 1
    for old_num in seen_in_text:
        old_to_new[old_num] = current_new_id
        new_to_old[current_new_id] = old_num
        current_new_id += 1
        
    # 4. Find what actually changes
    changes = []
    for old_num in old_to_new:
        new_num = old_to_new[old_num]
        if old_num != new_num:
            changes.append({
                'old': old_num,
                'new': new_num,
                'content': bibliography.get(old_num, "UNKNOWN")[:80] + "..."
            })
            
    # Sort by new number
    changes.sort(key=lambda x: x['new'])
    
    # 5. Output Report
    print("## KẾ HOẠCH ĐÁNH LẠI SỐ TÀI LIỆU THAM KHẢO (RENUMBERING PLAN)")
    print(f"Tổng số tài liệu: {len(seen_in_text)}")
    print(f"Số tài liệu cần thay đổi vị trí: {len(changes)}\n")
    
    if len(changes) == 0:
        print("Không có tài liệu nào bị sai thứ tự. Mọi thứ đã hoàn hảo!")
        return

    print("| TLTK Mới | TLTK Cũ | Thay đổi | Trích dẫn (rút gọn) |")
    print("|---|---|---|---|")
    for c in changes:
        print(f"| **[{c['new']}]** | [{c['old']}] | Cũ {c['old']} → Mới {c['new']} | {c['content']} |")

if __name__ == "__main__":
    plan_renumber()
