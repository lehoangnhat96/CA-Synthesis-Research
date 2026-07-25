import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

docx_path = "D:/1 Master's Ana Chem/1 Master's thesis/Carbon Aerogel/Thesi grad/DCLV_CA_02.07 Fe-N CA_Renumbered.docx"
doc = Document(docx_path)

start = -1
end = -1
for i, p in enumerate(doc.paragraphs):
    if 'TÀI LIỆU THAM KHẢO' in p.text.upper():
        start = i
    if 'PHỤ LỤC' in p.text.upper() and start != -1 and i > start:
        end = i
        break

print(f"Biblio range: {start} -> {end} ({end-start-1} entries)")
for i in range(start, min(start+10, end)):
    p = doc.paragraphs[i]
    print(f"[{i}] {repr(p.text[:100])}")
print("...")
for i in range(max(end-3, start+10), end+1):
    p = doc.paragraphs[i]
    print(f"[{i}] {repr(p.text[:100])}")
